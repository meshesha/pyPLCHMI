"""Faceplates: one template screen for many identical devices, opened as pop-up windows."""

import io

import pytest

from app.dashboards.layout import sanitize_layout
from app.faceplates.binding import (create_tags, instances_of_tag, parameters, point_map, resolve, scan,
                                    substitute_layout, suggest_prefix)
from app.models import PLC, AuditLog, Dashboard, DataPoint, Faceplate, FaceplateInstance
from app.plc.driver import get_manager
from tests.conftest import PASSWORD, add_point, login, make_user


def make_template(db, plc):
    """A Plane screen with the tags of the reference meter (Panel_01.MeterA.) and one fixed tag."""
    screen = Dashboard(name="Meter template", slug="meter-template", kind="image")
    db.session.add(screen)
    db.session.commit()
    volt = add_point(screen, plc, "Panel_01.MeterA.voltage", "DB1.DBD0:REAL", unit="V", log_enabled=True)
    power = add_point(screen, plc, "Panel_01.MeterA.power", "DB1.DBD4:REAL", unit="kW")
    reset = add_point(screen, plc, "Panel_01.MeterA.reset", "DB1.DBX8.0", writable=True)
    freq = add_point(screen, plc, "Site.frequency", "DB2.DBD0:REAL")
    screen.set_layout({"version": 1, "canvas": {"width": 400, "height": 300}, "elements": [
        {"id": "v", "type": "value", "point": volt.id, "x": 10, "y": 10, "w": 100, "h": 30},
        {"id": "p", "type": "value", "point": power.id, "x": 10, "y": 50, "w": 100, "h": 30},
        {"id": "b", "type": "button", "point": reset.id, "action": "toggle", "x": 10, "y": 90, "w": 100, "h": 30,
         "vis": "tag", "visPoint": power.id, "visOp": "gt", "visValue": "0"},
        {"id": "f", "type": "value", "point": freq.id, "x": 10, "y": 130, "w": 100, "h": 30},
    ]})
    fp = Faceplate(name="Electricity meter", screen=screen, prefix="Panel_01.MeterA.")
    db.session.add(fp)
    db.session.commit()
    return fp, {"voltage": volt, "power": power, "reset": reset, "freq": freq}


def meter_tags(plc, prefix, writable=True):
    """Tags of another meter."""
    return {"voltage": add_point(None, plc, prefix + "voltage", "DB10.DBD0:REAL"),
            "power": add_point(None, plc, prefix + "power", "DB10.DBD4:REAL"),
            "reset": add_point(None, plc, prefix + "reset", "DB10.DBX8.0", writable=writable)}


# ---------------------------------------------------------------------------
# Binding
# ---------------------------------------------------------------------------
def test_parameters_resolve_and_layout(db, plc):
    fp, ref = make_template(db, plc)
    assert [p.name for p in parameters(fp)] == ["power", "reset", "voltage"]
    assert suggest_prefix(fp.screen.datapoints) == "Panel_01.MeterA."
    tags = meter_tags(plc, "Panel_37.MeterB.")
    other = add_point(None, plc, "Spare.u1", "DB11.DBD0:REAL")
    inst = FaceplateInstance(faceplate=fp, name="Panel 37 B", prefix="panel_37.meterb.")    # case ignored
    db.session.add(inst)
    db.session.commit()
    found = resolve(inst)
    assert found["voltage"].id == tags["voltage"].id and found["reset"].id == tags["reset"].id
    inst.overrides = {"voltage": other.id}
    mapping, missing = point_map(inst)
    assert mapping[ref["voltage"].id] == other.id and mapping[ref["freq"].id] == ref["freq"].id and missing == []
    layout = substitute_layout(fp.screen.layout_data, mapping)
    by_id = {e["id"]: e for e in layout["elements"]}
    assert by_id["p"]["point"] == tags["power"].id and by_id["b"]["visPoint"] == tags["power"].id
    assert fp.screen.layout_data["elements"][0]["point"] == ref["voltage"].id      # template unchanged

    lonely = FaceplateInstance(faceplate=fp, name="Empty", prefix="Panel_99.X.")
    db.session.add(lonely)
    db.session.commit()
    mapping, missing = point_map(lonely)
    assert set(missing) == {"power", "reset", "voltage"}
    layout = substitute_layout(fp.screen.layout_data, mapping)
    button = next(e for e in layout["elements"] if e["id"] == "b")
    assert button["point"] is None and button["vis"] == "always"                  # no tag: visible, no value

    assert [i.name for i in instances_of_tag(tags["power"])] == ["Panel 37 B"]
    assert [i.name for i in instances_of_tag(other)] == ["Panel 37 B"]
    assert instances_of_tag(tags["voltage"]) == []                                 # replaced by the override


def test_scan_finds_devices(db, plc):
    fp, _ref = make_template(db, plc)
    meter_tags(plc, "Panel_02.MeterA.")
    add_point(None, plc, "Panel_03.MeterA.power", "DB12.DBD4:REAL")
    db.session.add(FaceplateInstance(faceplate=fp, name="P2", prefix="Panel_02.MeterA."))
    db.session.commit()
    found = {c["prefix"]: c for c in scan(fp)}
    assert "Panel_01.MeterA." not in found and "Panel_02.MeterA." not in found     # reference / existing
    assert found["Panel_03.MeterA."]["found"] == 1 and found["Panel_03.MeterA."]["total"] == 3
    assert found["Panel_03.MeterA."]["name"] == "Panel 03 · MeterA"


def test_create_tags_copies_reference(db, plc):
    fp, ref = make_template(db, plc)
    other_s7 = PLC(name="PLC2", ip_address="10.0.0.2", protocol="s7")
    modbus = PLC(name="MB", ip_address="10.0.0.3", protocol="modbus_tcp", tcp_port=502)
    db.session.add_all([other_s7, modbus])
    db.session.commit()
    created, problems = create_tags(fp, "Panel_05.MeterA.", other_s7)
    assert problems == [] and len(created) == 3
    db.session.commit()
    copy = DataPoint.query.filter_by(label="Panel_05.MeterA.voltage").one()
    assert copy.plc_id == other_s7.id and copy.address == ref["voltage"].address and copy.unit == "V"
    assert copy.log_enabled and not copy.screens
    assert create_tags(fp, "Panel_05.MeterA.", other_s7) == ([], [])                 # nothing missing
    created, problems = create_tags(fp, "Panel_06.MeterA.", modbus)
    assert created == [] and "another protocol" in problems[0]
    assert create_tags(fp, "6 bad", None)[1]


def test_layout_faceplate_reference_and_button():
    layout = sanitize_layout({"elements": [
        {"type": "button", "action": "faceplate", "faceplate": 5, "x": 0, "y": 0},
        {"type": "icon", "faceplate": 9, "x": 0, "y": 0},
    ]}, point_ids=[], asset_ids=[], dashboard_ids=[], faceplate_ids=[5])
    assert layout["elements"][0]["action"] == "faceplate" and layout["elements"][0]["faceplate"] == 5
    assert layout["elements"][1]["faceplate"] is None                               # unknown device dropped


# ---------------------------------------------------------------------------
# Admin pages
# ---------------------------------------------------------------------------
def test_type_and_device_pages(client, login_as, db, plc):
    login_as("admin")
    fp, ref = make_template(db, plc)
    assert client.get("/faceplates/").status_code == 200
    preview = client.get(f"/faceplates/api/preview?screen={fp.screen_id}&prefix=Panel_01.MeterA.").json
    assert [p["name"] for p in preview["parameters"]] == ["power", "reset", "voltage"]
    assert preview["fixed"] == ["Site.frequency"] and preview["suggested"] == "Panel_01.MeterA."

    resp = client.post("/faceplates/new", data={"name": "Bad", "screen_id": fp.screen_id, "prefix": "Nothing."})
    assert resp.status_code == 200 and "No tag of the screen starts with" in resp.get_data(as_text=True)
    resp = client.post("/faceplates/new", data={"name": "electricity METER", "screen_id": fp.screen_id,
                                                 "prefix": "Panel_01.MeterA."})
    assert "already exists" in resp.get_data(as_text=True)

    # A new device whose tags do not exist yet: create them as copies.
    resp = client.post(f"/faceplates/{fp.id}/instances/new", data={
        "name": "Panel 07 · A", "prefix": "Panel_07.MeterA.", "create_tags": "y", "plc_id": "0", "db_shift": "6"})
    assert resp.status_code == 302
    inst = FaceplateInstance.query.filter_by(name="Panel 07 · A").one()
    assert point_map(inst)[1] == [] and DataPoint.query.filter_by(label="Panel_07.MeterA.reset").one().writable
    page = client.get(f"/faceplates/{fp.id}").get_data(as_text=True)
    assert "Panel 07 · A" in page and "3 / 3" in page

    # Hand binding by tag name; an unknown name is refused.
    spare = add_point(None, plc, "Spare.u1", "DB11.DBD0:REAL")
    url = f"/faceplates/instances/{inst.id}/edit"
    resp = client.post(url, data={"name": "Panel 07 · A", "prefix": "Panel_07.MeterA.", "bind_voltage": "nope"})
    assert resp.status_code == 200 and "no tag named nope" in resp.get_data(as_text=True)
    resp = client.post(url, data={"name": "Panel 07 · A", "prefix": "Panel_07.MeterA.", "bind_voltage": "spare.U1"})
    assert resp.status_code == 302 and db.session.get(FaceplateInstance, inst.id).overrides == {"voltage": spare.id}

    # Find devices in the tags, and create them.
    meter_tags(plc, "Panel_08.MeterA.")
    cands = client.get(f"/faceplates/{fp.id}/scan").json["candidates"]
    assert [c["prefix"] for c in cands] == ["Panel_08.MeterA."]
    resp = client.post(f"/faceplates/{fp.id}/scan", json={"items": [{"prefix": "Panel_08.MeterA.", "name": ""}]})
    assert resp.json["created"] == 1 and FaceplateInstance.query.count() == 2

    # Excel export → import (update one, add one).
    xlsx = client.get(f"/faceplates/{fp.id}/instances.xlsx")
    assert xlsx.status_code == 200 and xlsx.data[:2] == b"PK"
    from openpyxl import Workbook, load_workbook  # noqa: PLC0415

    wb = load_workbook(io.BytesIO(xlsx.data))
    ws = wb.worksheets[0]
    assert [c.value for c in ws[1]][:4] == ["Name", "Prefix", "Description", "Bindings"]
    assert any(row[3] == "voltage=Spare.u1" for row in ws.iter_rows(min_row=2, values_only=True))
    ws.append(["Panel 09", "Panel_09.MeterA.", "new", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = client.post(f"/faceplates/{fp.id}/instances/import", data={"file": (buf, "d.xlsx")},
                       content_type="multipart/form-data")
    assert resp.status_code == 302 and FaceplateInstance.query.count() == 3
    bad = Workbook()
    bad.active.append(["Name", "Prefix"])
    bad.active.append(["X", "1 bad"])
    buf = io.BytesIO()
    bad.save(buf)
    buf.seek(0)
    client.post(f"/faceplates/{fp.id}/instances/import", data={"file": (buf, "d.xlsx")},
                content_type="multipart/form-data")
    assert FaceplateInstance.query.count() == 3                                      # nothing imported

    resp = client.post(f"/faceplates/{fp.id}/delete")
    assert resp.status_code == 302 and FaceplateInstance.query.count() == 0
    assert DataPoint.query.filter_by(label="Panel_07.MeterA.power").count() == 1    # tags kept


def test_designer_offers_devices(client, login_as, db, plc):
    login_as("admin")
    fp, _ref = make_template(db, plc)
    db.session.add(FaceplateInstance(faceplate=fp, name="P1", prefix="Panel_01.MeterA."))
    db.session.commit()
    page = client.get(f"/screens/{fp.screen_id}/designer").get_data(as_text=True)
    assert ("Electricity meter – P1" in page or "Electricity meter \\u2013 P1" in page) and '"faceplate"' in page


# ---------------------------------------------------------------------------
# Run time
# ---------------------------------------------------------------------------
def test_window_api_values_write_and_details(client, db, plc, users):
    fp, ref = make_template(db, plc)
    tags = meter_tags(plc, "Panel_37.MeterB.")
    inst = FaceplateInstance(faceplate=fp, name="Panel 37 B", prefix="Panel_37.MeterB.")
    db.session.add(inst)
    db.session.commit()
    get_manager().write(plc, tags["power"].tag_address, 12.5)
    make_user("op", "Operator")
    login(client, "op")

    data = client.get(f"/faceplates/api/instances/{inst.id}").json
    assert data["title"] == "Electricity meter – Panel 37 B" and data["missing"] == []
    points = {e["id"]: e["point"] for e in data["layout"]["elements"]}
    assert points == {"v": tags["voltage"].id, "p": tags["power"].id, "b": tags["reset"].id, "f": ref["freq"].id}
    assert data["points"][str(tags["reset"].id)]["may_write"] is True
    assert str(tags["voltage"].id) in data["menu_points"]
    values = client.get(f"/faceplates/api/instances/{inst.id}/values").json["values"]
    assert values[str(tags["power"].id)]["value"] == pytest.approx(12.5)

    # The tags are on no screen, but the window allows writing and Details.
    assert not tags["reset"].screens
    resp = client.post(f"/api/datapoints/{tags['reset'].id}/write", json={"value": True})
    assert resp.status_code == 200, resp.get_data(as_text=True)
    assert AuditLog.query.filter_by(action="write", object_id=tags["reset"].id).count() == 1
    details = client.get(f"/api/datapoints/{tags['power'].id}/details").json
    assert details["devices"] == ["Electricity meter – Panel 37 B"]
    unrelated = add_point(None, plc, "Other.tag", "DB20.DBW0:INT", writable=True)
    assert client.post(f"/api/datapoints/{unrelated.id}/write", json={"value": 1}).status_code == 403

    page = client.get("/faceplates/devices").get_data(as_text=True)
    assert 'data-faceplate="%d"' % inst.id in page and "hmi-faceplate.js" in page
    assert 'id="nav-devices"' in page


def test_window_needs_access_to_template(client, db, plc, users):
    fp, _ref = make_template(db, plc)
    tags = meter_tags(plc, "Panel_37.MeterB.")
    inst = FaceplateInstance(faceplate=fp, name="P37", prefix="Panel_37.MeterB.")
    db.session.add(inst)
    db.session.commit()
    login(client, "guest", PASSWORD)                      # Guest: public screens only, template is not public
    assert client.get(f"/faceplates/api/instances/{inst.id}").status_code == 403
    assert client.get(f"/faceplates/api/instances/{inst.id}/values").status_code == 403
    assert client.get(f"/api/datapoints/{tags['power'].id}/details").status_code == 403
    assert f'data-faceplate="{inst.id}"' not in client.get("/faceplates/devices").get_data(as_text=True)
    client.post("/auth/logout")
    assert client.get(f"/faceplates/api/instances/{inst.id}").status_code == 401
    fp.screen.is_public = True
    db.session.commit()
    login(client, "guest", PASSWORD)
    assert client.get(f"/faceplates/api/instances/{inst.id}").status_code == 200


def test_screen_button_opens_device(client, login_as, db, plc):
    login_as("admin")
    fp, _ref = make_template(db, plc)
    inst = FaceplateInstance(faceplate=fp, name="P1", prefix="Panel_01.MeterA.")
    screen = Dashboard(name="Panels", slug="panels", kind="image")
    db.session.add_all([inst, screen])
    db.session.commit()
    resp = client.put(f"/api/dashboards/{screen.id}/layout", json={"elements": [
        {"type": "button", "action": "faceplate", "faceplate": inst.id, "text": "Meter A", "x": 0, "y": 0}]})
    assert resp.status_code == 200 and resp.json["layout"]["elements"][0]["faceplate"] == inst.id
    page = client.get("/screens/panels").get_data(as_text=True)
    assert "hmi-faceplate.js" in page and "faceplate-config" in page


def test_create_tags_with_address_shift(db, plc):
    fp, ref = make_template(db, plc)
    created, problems = create_tags(fp, "Panel_02.MeterA.")                         # same PLC, no shift
    assert created == [] and "same addresses" in problems[0]
    created, problems = create_tags(fp, "Panel_02.MeterA.", None, db_shift=1)
    assert problems == []
    db.session.commit()
    addresses = {dp.label: dp.address for dp in created}
    assert addresses == {"Panel_02.MeterA.voltage": "DB2.DBD0", "Panel_02.MeterA.power": "DB2.DBD4",
                         "Panel_02.MeterA.reset": "DB2.DBX8.0"}
    created, problems = create_tags(fp, "Panel_03.MeterA.", None, start_shift=20)  # 20 bytes per meter
    db.session.commit()
    assert {dp.label: dp.address for dp in created}["Panel_03.MeterA.reset"] == "DB1.DBX28.0"
    assert get_manager().read_one(plc, created[0].tag_address).ok
    assert "negative" in create_tags(fp, "Panel_04.MeterA.", None, start_shift=-100)[1][0]


def test_import_creates_devices_and_tags(client, login_as, db, plc):
    from openpyxl import Workbook  # noqa: PLC0415

    login_as("admin")
    fp, _ref = make_template(db, plc)
    wb = Workbook()
    wb.active.append(["Name", "Prefix", "Create tags", "PLC", "DB shift", "Address shift"])
    for n in range(2, 6):
        wb.active.append([f"Panel {n}", f"Panel_{n:02d}.MeterA.", "yes", "", n - 1, ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = client.post(f"/faceplates/{fp.id}/instances/import", data={"file": (buf, "d.xlsx")},
                       content_type="multipart/form-data", follow_redirects=True)
    assert "4 device(s) added" in resp.get_data(as_text=True) and "12 tag(s) created" in resp.get_data(as_text=True)
    assert DataPoint.query.filter_by(label="Panel_05.MeterA.voltage").one().address == "DB5.DBD0"
    assert all(point_map(i)[1] == [] for i in FaceplateInstance.query.all())

    bad = Workbook()                                         # same PLC without a shift: nothing at all
    bad.active.append(["Name", "Prefix", "Create tags"])
    bad.active.append(["Panel 9", "Panel_09.MeterA.", "yes"])
    buf = io.BytesIO()
    bad.save(buf)
    buf.seek(0)
    resp = client.post(f"/faceplates/{fp.id}/instances/import", data={"file": (buf, "d.xlsx")},
                       content_type="multipart/form-data", follow_redirects=True)
    assert "Nothing imported" in resp.get_data(as_text=True)
    assert FaceplateInstance.query.count() == 4 and DataPoint.query.filter(DataPoint.label.like("Panel_09%")).count() == 0
