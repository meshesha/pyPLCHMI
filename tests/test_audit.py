"""Tests for the audit log (automatic and explicit entries, page, permissions)."""

import io

from app.audit.recorder import paused
from app.models import PLC, AuditLog, Role
from tests.conftest import PASSWORD


def entries(action=None, object_type=None):
    """Audit rows, oldest first."""
    q = AuditLog.query
    if action:
        q = q.filter_by(action=action)
    if object_type:
        q = q.filter_by(object_type=object_type)
    return q.order_by(AuditLog.id).all()


def change(entry, field):
    """The change of one field of an entry."""
    return next(c for c in entry.changes if c["field"] == field)


def test_login_logout_and_failed_login(client, users):
    """Log in, failed log in and log out are recorded with user and IP."""
    client.post("/auth/login", data={"username": "admin", "password": "wrong"})
    client.post("/auth/login", data={"username": "admin", "password": PASSWORD})
    client.post("/auth/logout")
    failed, login = entries("login_failed")[0], entries("login")[0]
    assert failed.username == "admin" and failed.user_id is None and failed.ip == "127.0.0.1"
    assert login.user_id == users["admin"].id and entries("logout")[0].username == "admin"


def test_create_update_delete_with_before_and_after(client, login_as, users):
    """Changes of a PLC are recorded with the values before and after."""
    login_as("admin")
    client.post("/admin/plcs/new", data={"name": "Boiler", "protocol": "s7", "ip_address": "10.0.0.9", "tcp_port": 102,
                                         "rack": 0, "slot": 1, "enabled": "y", "offline_delay": 5, "timeout": 3,
                                         "description": "first"})
    plc = PLC.query.filter_by(name="Boiler").first()
    assert plc is not None, "PLC form did not save"
    created = entries("create", "PLC")[-1]
    assert created.object_label == "Boiler" and created.username == "admin"
    assert change(created, "ip_address") == {"field": "ip_address", "before": None, "after": "10.0.0.9"}
    client.post(f"/admin/plcs/{plc.id}/edit", data={"name": "Boiler", "protocol": "s7", "ip_address": "10.0.0.10",
                                                   "tcp_port": 102, "rack": 0, "slot": 1, "enabled": "y",
                                                   "offline_delay": 5, "timeout": 3, "description": "second"})
    updated = entries("update", "PLC")[-1]
    assert change(updated, "ip_address") == {"field": "ip_address", "before": "10.0.0.9", "after": "10.0.0.10"}
    assert change(updated, "description")["before"] == "first"
    assert all(c["field"] != "name" for c in updated.changes)  # unchanged fields are not listed
    client.post(f"/admin/plcs/{plc.id}/delete")
    deleted = entries("delete", "PLC")[-1]
    assert deleted.object_id == plc.id and change(deleted, "ip_address")["before"] == "10.0.0.10"


def test_passwords_are_never_logged(client, login_as, users):
    """Password changes are recorded without the hash."""
    login_as("admin")
    client.post("/auth/profile", data={"current_password": PASSWORD, "new_password": "another-pass-9",
                                       "confirm": "another-pass-9"})
    entry = entries("update", "User")[-1]
    assert change(entry, "password_hash") == {"field": "password_hash", "before": "•••", "after": "•••"}
    assert "$2b$" not in entry.changes_json


def test_many_to_many_and_foreign_keys(client, login_as, users, db):
    """Role permissions are listed by name; foreign keys show the referenced name."""
    login_as("admin")
    viewer = Role.query.filter_by(name="Viewer").one()
    viewer.set_permissions(["view_dashboards"])
    with client.application.test_request_context("/"):
        db.session.commit()
    entry = entries("update", "Role")[-1]
    perms = change(entry, "permissions")
    assert "view_history" in perms["before"] and perms["after"] == "view_dashboards"
    client.post(f"/admin/users/{users['viewer'].id}/edit", data={"username": "viewer", "role_id": viewer.id - 1,
                                                                "active": "y"})
    role_change = change(entries("update", "User")[-1], "role_id")
    assert "(Viewer)" in role_change["before"] and "(Admin)" in role_change["after"]


def test_value_write_recorded(client, login_as, dashboards):
    """Writing a value to a PLC records the value before and after."""
    login_as("admin")
    sp = dashboards["points"]["setpoint"]
    client.post(f"/api/datapoints/{sp.id}/write", json={"value": 42.5})
    entry = entries("write")[-1]
    assert entry.object_type == "Value" and "Setpoint" in entry.object_label
    assert change(entry, "value")["after"] == 42.5


def test_background_and_paused_changes_not_recorded(app, db, plc):
    """Changes outside a request (CLI, background threads) or while paused are not recorded."""
    plc.description = "changed by the system"
    db.session.commit()
    with app.test_request_context("/"):
        with paused():
            plc.description = "paused"
            db.session.commit()
    assert entries() == []


def test_audit_page_filters_export_and_permission(client, login_as, users):
    """Only roles with view_audit see the page; filters and Excel export work."""
    login_as("viewer")
    assert client.get("/admin/audit/").status_code == 403
    assert 'id="nav-audit"' not in client.get("/screens/").get_data(as_text=True)
    client.post("/auth/logout")
    login_as("admin")
    assert 'id="nav-audit"' in client.get("/screens/").get_data(as_text=True)
    html = client.get("/admin/audit/?action=login").get_data(as_text=True)
    assert 'id="audit-table"' in html and "audit-badge-login" in html and "audit-badge-logout" not in html
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(io.BytesIO(client.get("/admin/audit/export.xlsx").data))
    rows = list(wb["Audit log"].iter_rows(values_only=True))
    assert rows[0][:4] == ("Date & time", "User", "IP", "Action") and any(r[3] == "Logged in" for r in rows[1:])


def test_library_upload_recorded(client, login_as, plc):
    """Uploading and attaching library files is recorded."""
    login_as("admin")
    fid = client.post("/library/api/files", data={"file": (io.BytesIO(b"%PDF-1.4 x"), "m.pdf")},
                      content_type="multipart/form-data").get_json()["file"]["id"]
    client.post(f"/library/api/plcs/{plc.id}/files", json={"file_id": fid})
    assert entries("create", "Library file")[-1].object_label == "m"
    files = change(entries("update", "PLC")[-1], "files")
    assert files == {"field": "files", "before": "", "after": "m"}
    assert not entries("update", "Library file")  # the same change is not recorded twice
