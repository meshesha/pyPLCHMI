"""Tests for shared tags (screen links), the existing-tag picker, tag copy, dummy PLCs and Excel export/import."""

import io
import re

import pytest
from openpyxl import Workbook, load_workbook

from app.models import PLC, Alarm, Dashboard, DataPoint, ScreenTag
from app.plc.dummy import DummyConnection
from tests.conftest import add_point


@pytest.fixture(autouse=True)
def _reset_dummy():
    """Start every test with empty dummy PLC memory."""
    DummyConnection.reset()
    yield
    DummyConnection.reset()


def tag_form(plc, **over):
    """Valid tag form data.

    Args:
        plc: PLC of the tag.
        **over: Overrides.

    Returns:
        Dict for POST.
    """
    data = {"label": "New tag", "plc_id": plc.id, "address": "DB9.DBW0", "data_type": "INT", "widget": "value",
            "width": 1, "scale": 1, "offset": 0, "decimals": 0, "true_text": "ON", "false_text": "OFF",
            "string_length": 20}
    data.update(over)
    return data


@pytest.fixture()
def dummy(db):
    """A dummy PLC.

    Returns:
        :class:`PLC`.
    """
    plc = PLC(name="Sim", ip_address="", protocol="dummy")
    db.session.add(plc)
    db.session.commit()
    return plc


# ---------------------------------------------------------------- shared tags
def test_tag_on_two_screens(client, login_as, db, dashboards):
    """One tag can be shown on several screens; each keeps its own order."""
    temp, public, private = dashboards["points"]["temp"], dashboards["public"], dashboards["private"]
    public.add_tag(temp)
    db.session.commit()
    assert {s.slug for s in temp.screens} == {"public", "private"}
    assert public.datapoints[-1].id == temp.id and private.datapoints[0].id == temp.id
    login_as("viewer")
    data = client.get(f"/api/dashboards/{public.id}/values").get_json()
    assert str(temp.id) in data["values"]


def test_remove_from_screen_keeps_tag(client, login_as, db, dashboards):
    """'Remove from screen' only removes the link."""
    login_as("admin")
    temp, private = dashboards["points"]["temp"], dashboards["private"]
    html = client.get(f"/screens/{private.id}/edit").get_data(as_text=True)
    assert "Remove from screen" in html and 'id="add-existing-btn"' in html and "tag_picker.js" in html
    client.post(f"/screens/{private.id}/points/{temp.id}/delete")
    db.session.expire_all()
    assert db.session.get(DataPoint, temp.id) is not None
    assert not db.session.get(Dashboard, private.id).has_tag(temp.id)


def test_deleting_tag_removes_links(db, dashboards):
    """Deleting a tag deletes its screen links."""
    temp = dashboards["points"]["temp"]
    dashboards["public"].add_tag(temp)
    db.session.commit()
    db.session.delete(temp)
    db.session.commit()
    assert ScreenTag.query.filter_by(datapoint_id=temp.id).count() == 0


def test_add_existing_tags_to_grid_screen(client, login_as, db, dashboards, plc):
    """The grid editor links chosen tags at the end; tags already there are skipped."""
    login_as("admin")
    loose = add_point(None, plc, "Loose", "MW40:INT")
    public = dashboards["public"]
    resp = client.post(f"/screens/{public.id}/points/add",
                       data={"datapoint_id": [loose.id, dashboards["points"]["pub"].id, "x"]}, follow_redirects=True)
    assert "Added to this screen: Loose" in resp.get_data(as_text=True)
    db.session.expire_all()
    screen = db.session.get(Dashboard, public.id)
    assert [dp.label for dp in screen.datapoints] == ["Public value", "Loose"]


def test_write_allowed_through_any_screen(client, login_as, db, dashboards, users, plc):
    """A tag may be written from any screen the user can open."""
    from app.models import Perm, Permission  # noqa: PLC0415
    users["guest"].role.permissions.append(Permission.query.filter_by(code=Perm.WRITE_VALUES).one())
    sp = dashboards["points"]["setpoint"]
    login_as("guest")
    assert client.post(f"/api/datapoints/{sp.id}/write", json={"value": 5}).status_code == 403
    dashboards["public"].add_tag(sp)
    db.session.commit()
    assert client.post(f"/api/datapoints/{sp.id}/write", json={"value": 5}).status_code == 200


# ---------------------------------------------------------------- picker API
def test_tag_search_api(client, login_as, dashboards, plc):
    """The picker searches label/address/PLC and marks tags already on the screen."""
    login_as("admin")
    private = dashboards["private"]
    data = client.get(f"/api/tags/search?q=DB1.DBD&screen={private.id}").get_json()
    labels = {t["label"]: t for t in data["tags"]}
    assert set(labels) == {"Temp", "Setpoint"} and labels["Temp"]["on_screen"] is True
    data = client.get(f"/api/tags/search?q=public&screen={private.id}").get_json()
    assert [t["label"] for t in data["tags"]] == [] or all(not t["on_screen"] for t in data["tags"])
    data = client.get("/api/tags/search?q=PLC1&limit=2").get_json()
    assert len(data["tags"]) == 2 and data["more"] is True


def test_tag_search_requires_screen_manager(client, login_as, dashboards):
    """Viewers cannot use the picker API."""
    login_as("viewer")
    assert client.get("/api/tags/search").status_code == 403


def test_designer_link_and_unlink_api(client, login_as, db, plc):
    """The designer links existing tags; its layout may then use them."""
    login_as("admin")
    screen = Dashboard(name="Plant", slug="plant", kind="image")
    db.session.add(screen)
    db.session.commit()
    loose = add_point(None, plc, "Level", "MW60:INT")
    html = client.get(f"/screens/{screen.id}/designer").get_data(as_text=True)
    assert 'id="tag-picker"' in html and "tag_picker.js" in html
    body = client.post(f"/api/dashboards/{screen.id}/tags", json={"datapoint_id": loose.id}).get_json()
    assert body["id"] == loose.id and body["label"] == "Level"
    layout = {"version": 1, "canvas": {"width": 800, "height": 600, "background": "#ffffff"},
              "elements": [{"id": "v1", "type": "value", "x": 10, "y": 10, "w": 100, "h": 30, "point": loose.id}]}
    saved = client.put(f"/api/dashboards/{screen.id}/layout", json=layout).get_json()
    assert saved["layout"]["elements"][0]["point"] == loose.id
    assert client.post(f"/api/dashboards/{screen.id}/tags", json={}).status_code == 400
    assert client.delete(f"/api/dashboards/{screen.id}/tags/{loose.id}").get_json() == {"removed": True}
    assert client.delete(f"/api/dashboards/{screen.id}/tags/{loose.id}").get_json() == {"removed": False}


def test_new_tag_from_screen_goes_to_tag_table(client, login_as, db, dashboards, plc):
    """Tags created in the screen editor appear in the tag table."""
    login_as("admin")
    private = dashboards["private"]
    client.post(f"/screens/{private.id}/points/new", data=tag_form(plc, label="Fresh"))
    tag = DataPoint.query.filter_by(label="Fresh").one()
    assert [s.id for s in tag.screens] == [private.id]
    assert "Fresh" in client.get("/tags/").get_data(as_text=True)


# ---------------------------------------------------------------- copy
def test_copy_button_prefills_form(client, login_as, db, dashboards):
    """The Copy button opens the new-tag form filled with the tag's settings."""
    login_as("admin")
    speed = dashboards["points"]["speed"]
    html = client.get("/tags/").get_data(as_text=True)
    assert f'/tags/new?copy={speed.id}' in html
    html = client.get(f"/tags/new?copy={speed.id}").get_data(as_text=True)
    assert 'value="Speed (copy)"' in html and 'value="DB1.DBW8"' in html and "Copy tag" in html
    assert re.search(r'<option selected value="INT">|<option value="INT" selected', html)
    assert 'value="0.1"' in html                      # scale copied
    client.post(f"/tags/new?copy={speed.id}", data=tag_form(speed.plc, label="Speed 2", address="DB1.DBW10",
                                                             scale=0.1, decimals=1, writable="y"))
    copy = DataPoint.query.filter_by(label="Speed 2").one()
    assert copy.id != speed.id and copy.address == "DB1.DBW10" and copy.scale == 0.1
    assert db.session.get(DataPoint, speed.id).label == "Speed"


# ---------------------------------------------------------------- dummy PLC
def test_dummy_plc_form_needs_no_ip(client, login_as, db):
    """A dummy PLC is saved without IP address, gets no communication alarm, and tests OK."""
    login_as("admin")
    resp = client.post("/admin/plcs/new", data={
        "name": "Internal", "protocol": "dummy", "ip_address": "", "timeout": 3, "offline_delay": 10,
        "modbus_library": "pymodbus", "data_order": "ABCD", "bacnet_priority": 8, "baudrate": 9600, "parity": "N",
        "bytesize": 8, "stopbits": 1, "comm_alarm": "y", "comm_severity": "high", "enabled": "y"})
    assert resp.status_code == 302, resp.get_data(as_text=True)
    plc = PLC.query.filter_by(name="Internal").one()
    assert plc.is_dummy and plc.endpoint == "no device" and plc.comm_alarm is None
    resp = client.post(f"/admin/plcs/{plc.id}/test", follow_redirects=True)
    assert "Dummy PLC" in resp.get_data(as_text=True)
    html = client.get("/admin/plcs").get_data(as_text=True)
    assert "internal values" in html and "Dummy (no device" in html


def test_dummy_tags_without_address(client, login_as, db, dummy):
    """Dummy tags need no address; each gets the next internal slot."""
    login_as("admin")
    assert client.post("/tags/new", data=tag_form(dummy, label="Mode", address="", data_type="INT",
                                                  writable="y")).status_code == 302
    assert client.post("/tags/new", data=tag_form(dummy, label="Run", address="ignored", data_type="BOOL",
                                                  widget="indicator")).status_code == 302
    mode, run = DataPoint.query.filter_by(label="Mode").one(), DataPoint.query.filter_by(label="Run").one()
    assert (mode.address, run.address) == ("D1", "D2") and run.data_type == "BOOL"
    html = client.post("/tags/new", data=tag_form(dummy, label="T", address="", data_type="TIME")).get_data(as_text=True)
    assert "not available for dummy tags" in html
    # editing keeps the slot
    client.post(f"/tags/{mode.id}/edit", data=tag_form(dummy, label="Mode 2", address="", writable="y"))
    db.session.refresh(mode)
    assert mode.address == "D1" and mode.label == "Mode 2"
    edit_html = client.get(f"/tags/{mode.id}/edit").get_data(as_text=True)
    assert re.search(r'id="address"[^>]*value=""|value=""[^>]*id="address"', edit_html)


def test_dummy_tag_values_read_and_write(client, login_as, db, dummy, dashboards):
    """Dummy values start at 0 and keep what operators write."""
    login_as("admin")
    tag = DataPoint(plc=dummy, label="Counter", writable=True)
    from app.plc.dummy import DummyAddress  # noqa: PLC0415
    tag.set_address(DummyAddress(start=1, data_type="REAL"))
    db.session.add(tag)
    dashboards["public"].add_tag(tag)
    db.session.commit()
    url = f"/api/dashboards/{dashboards['public'].id}/values"
    assert client.get(url).get_json()["values"][str(tag.id)]["value"] == 0.0
    assert client.post(f"/api/datapoints/{tag.id}/write", json={"value": 12.5}).get_json()["value"] == 12.5
    v = client.get(url).get_json()
    assert v["values"][str(tag.id)]["value"] == 12.5 and v["plcs"][str(dummy.id)]["state"] == "online"


def test_dummy_tag_moves_to_real_plc(client, login_as, db, dummy, plc):
    """A dummy tag can later be connected to a real PLC (address then required)."""
    login_as("admin")
    client.post("/tags/new", data=tag_form(dummy, label="Later", address=""))
    tag = DataPoint.query.filter_by(label="Later").one()
    html = client.post(f"/tags/{tag.id}/edit", data=tag_form(plc, label="Later", address="")).get_data(as_text=True)
    assert "Enter the address" in html
    client.post(f"/tags/{tag.id}/edit", data=tag_form(plc, label="Later", address="DB4.DBW2"))
    db.session.refresh(tag)
    assert tag.plc_id == plc.id and tag.address == "DB4.DBW2"
    # …and back to the dummy PLC: a fresh slot
    client.post(f"/tags/{tag.id}/edit", data=tag_form(dummy, label="Later", address="DB4.DBW2"))
    db.session.refresh(tag)
    assert tag.plc_id == dummy.id and tag.address == "D1"


def test_designer_api_creates_dummy_tag(client, login_as, db, dummy):
    """The designer's New-tag dialog works without an address for dummy PLCs."""
    login_as("admin")
    screen = Dashboard(name="Mimic", slug="mimic", kind="image")
    db.session.add(screen)
    db.session.commit()
    resp = client.post(f"/api/dashboards/{screen.id}/datapoints", json={"label": "Lamp", "plc_id": dummy.id,
                                                                         "data_type": "BOOL", "writable": True})
    assert resp.status_code == 201 and resp.get_json()["address"] == "D1"


# ---------------------------------------------------------------- Excel: tags
def _sheet(resp):
    """Rows of the first sheet of a downloaded workbook.

    Args:
        resp: Flask response.

    Returns:
        List of row tuples (header first).
    """
    assert resp.mimetype.endswith("spreadsheetml.sheet")
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames[1] == "Help"
    return list(wb.worksheets[0].iter_rows(values_only=True))


def _upload(rows, sheet="Tags"):
    """Build an .xlsx upload.

    Args:
        rows: Rows (header first).
        sheet: Sheet title.

    Returns:
        ``(BytesIO, filename)`` for the test client.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, "import.xlsx"


def test_tag_export_scopes(client, login_as, db, dashboards):
    """Export all, the current filter, or selected tags."""
    login_as("admin")
    rows = _sheet(client.get("/tags/export?scope=all"))
    assert rows[0][:4] == ("ID", "Label", "PLC", "Address") and len(rows) == 7
    temp = next(r for r in rows if r[1] == "Temp")
    assert temp[3] == "DB1.DBD0" and temp[4] == "REAL" and temp[17] == "Private"
    rows = _sheet(client.get("/tags/export?q=speed"))
    assert [r[1] for r in rows[1:]] == ["Speed"]
    ids = [dashboards["points"]["run"].id, dashboards["points"]["name"].id]
    rows = _sheet(client.post("/tags/export", data={"ids": ids}))
    assert sorted(r[1] for r in rows[1:]) == ["Name", "Run"]
    html = client.get("/tags/").get_data(as_text=True)
    assert 'id="export-selected"' in html and 'class="form-check-input row-select"' in html and 'id="import-modal"' in html


def test_tag_import_roundtrip_update_and_create(client, login_as, db, dashboards, plc, dummy):
    """Exported rows update tags by ID; rows without ID create tags (also dummy ones and screen links)."""
    login_as("admin")
    rows = [list(r) for r in _sheet(client.get("/tags/export?scope=all"))]
    header = rows[0]
    temp_row = next(r for r in rows if r[1] == "Temp")
    temp_row[1] = "Temperature"
    temp_row[6] = "K"
    new_real = [None, "Flow", "PLC1", "DB7.DBD0", "REAL", None, "m3/h", 1, 0, 1, None, None, "ON", "OFF", "yes",
                "value", 2, "Public, Private"]
    new_dummy = [None, "Mode", "Sim", None, "INT", None, "", 1, 0, 0, None, None, "ON", "OFF", "yes", "value", 1, ""]
    speed_row = next(r for r in rows if r[1] == "Speed")          # unchanged
    upload = _upload([header, temp_row, speed_row, new_real, new_dummy])
    html = client.post("/tags/import", data={"file": upload}, content_type="multipart/form-data").get_data(as_text=True)
    text = " ".join(re.sub(r"<[^>]+>", " ", html).split())
    assert "2 tag(s) created, 1 updated, 1 unchanged" in text, text[text.find("Imported"):][:900]
    db.session.expire_all()
    temp = db.session.get(DataPoint, dashboards["points"]["temp"].id)
    assert temp.label == "Temperature" and temp.unit == "K" and [s.slug for s in temp.screens] == ["private"]
    flow = DataPoint.query.filter_by(label="Flow").one()
    assert flow.writable and flow.width == 2 and {s.slug for s in flow.screens} == {"public", "private"}
    assert DataPoint.query.filter_by(label="Mode").one().address == "D1"


def test_tag_import_errors_save_nothing(client, login_as, db, dashboards):
    """Any invalid row aborts the whole import and is reported."""
    login_as("admin")
    header = ["Label", "PLC", "Address", "Data type", "Screens"]
    upload = _upload([header, ["Good", "PLC1", "DB1.DBW50", "INT", ""], ["Bad", "PLC1", "XYZ", "INT", ""],
                      ["NoPlc", "Nope", "DB1.DBW52", "INT", ""], ["Scr", "PLC1", "DB1.DBW54", "INT", "Missing"]])
    html = client.post("/tags/import", data={"file": upload}, content_type="multipart/form-data").get_data(as_text=True)
    assert "Nothing was saved" in html and "3 of 4 row(s) have errors" in html
    assert "PLC &#39;Nope&#39; does not exist" in html and "Unknown screen(s): Missing" in html
    assert DataPoint.query.filter_by(label="Good").count() == 0


def test_tag_import_check_only_and_bad_files(client, login_as, db, dashboards):
    """Check-only validates without saving; unreadable files are rejected."""
    login_as("admin")
    upload = _upload([["Label", "PLC", "Address", "Data type"], ["Extra", "PLC1", "DB1.DBW60", "INT"]])
    html = client.post("/tags/import", data={"file": upload, "check_only": "1"},
                       content_type="multipart/form-data").get_data(as_text=True)
    assert "Check only" in html and "would create" in html
    assert DataPoint.query.filter_by(label="Extra").count() == 0
    resp = client.post("/tags/import", data={"file": (io.BytesIO(b"not excel"), "x.xlsx")},
                       content_type="multipart/form-data", follow_redirects=True)
    assert "not an Excel .xlsx file" in resp.get_data(as_text=True)
    resp = client.post("/tags/import", data={"file": _upload([["Foo"], ["bar"]])},
                       content_type="multipart/form-data", follow_redirects=True)
    assert "Missing column(s): label, plc" in resp.get_data(as_text=True)
    resp = client.post("/tags/import", data={}, follow_redirects=True)
    assert "Choose an Excel file" in resp.get_data(as_text=True)


@pytest.mark.parametrize("who", ["viewer", "guest"])
def test_excel_admin_only(client, login_as, dashboards, who):
    """Only administrators can export or import."""
    login_as(who)
    for url in ("/tags/export?scope=all", "/alarms/export?scope=all"):
        assert client.get(url).status_code == 403
    assert client.post("/tags/import").status_code == 403
    assert client.post("/alarms/import").status_code == 403


# ---------------------------------------------------------------- Excel: alarms
@pytest.fixture()
def alarms(db, dashboards, plc):
    """A tag alarm and a communication alarm.

    Returns:
        Dict ``{"hot", "comm"}``.
    """
    hot = Alarm(name="Hot", description="Too hot", datapoint=dashboards["points"]["temp"], operator="ge",
                setpoint="80", severity="critical", screen=dashboards["private"])
    db.session.add(hot)
    comm = Alarm.sync_comm_alarm(plc, True, "low")
    db.session.commit()
    return {"hot": hot, "comm": comm}


def test_alarm_export(client, login_as, alarms):
    """Alarms export with trigger, tag, condition symbol, severity label and screen."""
    login_as("admin")
    rows = _sheet(client.get("/alarms/export?scope=all"))
    hot = next(r for r in rows if r[1] == "Hot")
    assert hot[3:9] == ("tag", alarms["hot"].datapoint_id, "Temp", "PLC1", "≥", "80")
    assert hot[10:13] == ("Critical", "Private", "yes")
    comm = next(r for r in rows if r[1] == "PLC1 communication lost")
    assert comm[3] == "comm" and comm[9] == "PLC1" and comm[10] == "Low"
    rows = _sheet(client.get("/alarms/export?kind=comm"))
    assert [r[1] for r in rows[1:]] == ["PLC1 communication lost"]
    rows = _sheet(client.post("/alarms/export", data={"ids": [alarms["hot"].id]}))
    assert [r[1] for r in rows[1:]] == ["Hot"]


def test_alarm_import(client, login_as, db, alarms, dashboards):
    """Alarm import updates by ID and creates new ones (tag by label + PLC, symbols, severity labels)."""
    login_as("admin")
    header = ["ID", "Name", "Description", "Trigger", "Tag ID", "Tag", "Tag PLC", "Condition", "Value", "PLC",
              "Severity", "Screen", "Enabled"]
    rows = [header,
            [alarms["hot"].id, "Hot", "Way too hot", "tag", None, "Temp", "PLC1", ">", "90", None, "High", "", "no"],
            [None, "Pump off", "", "tag", None, "Run", "PLC1", "=", "false", None, "Low", "Public", "yes"],
            [None, "Low speed", "", "tag", dashboards["points"]["speed"].id, "", "", "<=", "5", None, "low", "", ""]]
    html = client.post("/alarms/import", data={"file": _upload(rows, "Alarms")},
                       content_type="multipart/form-data").get_data(as_text=True)
    assert "2 alarm(s) created, 1 updated, 0 unchanged" in " ".join(html.split())
    db.session.expire_all()
    hot = db.session.get(Alarm, alarms["hot"].id)
    assert (hot.operator, hot.setpoint, hot.severity, hot.screen_id, hot.enabled, hot.description) == \
        ("gt", "90", "high", None, False, "Way too hot")
    pump = Alarm.query.filter_by(name="Pump off").one()
    assert pump.datapoint.label == "Run" and pump.operator == "eq" and pump.screen.slug == "public"
    assert Alarm.query.filter_by(name="Low speed").one().operator == "le"


def test_alarm_import_errors(client, login_as, db, alarms):
    """Bad conditions, unknown tags, duplicate names and comm duplicates are reported; nothing is saved."""
    login_as("admin")
    header = ["Name", "Trigger", "Tag", "Tag PLC", "Condition", "Value", "PLC", "Severity"]
    rows = [header,
            ["A", "tag", "Temp", "PLC1", "~", "1", None, "High"],
            ["B", "tag", "Nope", "PLC1", ">", "1", None, "High"],
            ["Hot", "tag", "Temp", "PLC1", ">", "1", None, "High"],
            ["C", "comm", None, None, None, None, "PLC1", "High"],
            ["D", "tag", "Run", "PLC1", ">", "true", None, "High"],
            ["E", "tag", "Temp", "PLC1", ">", "1", None, "Urgent"]]
    html = client.post("/alarms/import", data={"file": _upload(rows, "Alarms")},
                       content_type="multipart/form-data").get_data(as_text=True)
    assert "6 of 6 row(s) have errors" in html
    for text in ("Unknown condition", "Tag &#39;Nope&#39; on PLC &#39;PLC1&#39; not found", "already exists",
                 "already has a communication alarm", "BOOL", "Unknown severity"):
        assert text in html, text
    assert Alarm.query.count() == 2
