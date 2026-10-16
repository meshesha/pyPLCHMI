"""Tests for general (database) and local (browser window) dummy tags."""

import io
import json
import re

import pytest
from openpyxl import load_workbook

from app.alarms.engine import scan_alarms
from app.models import PLC, Alarm, Dashboard, DataPoint, DummyValue
from app.plc.dummy import DummyAddress, DummyConnection, next_slot
from app.plc.service import read_datapoints
from tests.test_tag_features import tag_form


@pytest.fixture(autouse=True)
def _reset_dummy():
    """Start every test with an empty in-memory copy."""
    DummyConnection.reset()
    yield
    DummyConnection.reset()


@pytest.fixture()
def dummy(db):
    """A dummy PLC.

    Returns:
        :class:`PLC`.
    """
    plc = PLC(name="Sim", ip_address="", protocol="dummy")
    db.session.add(plc)
    db.session.commit()
    return plc


def make_tag(db, plc, label, dtype="INT", slot=1, scope="general", screen=None, **kw):
    """Create a dummy tag.

    Args:
        db: Database fixture.
        plc: Dummy PLC.
        label: Label.
        dtype: Data type.
        slot: Slot number.
        scope: ``general`` or ``local``.
        screen: Optional screen to show it on.
        **kw: Further column values.

    Returns:
        The :class:`DataPoint`.
    """
    tag = DataPoint(plc=plc, label=label, writable=True, dummy_scope=scope, **kw)
    tag.set_address(DummyAddress(start=slot, data_type=dtype, string_length=10 if dtype == "STRING" else 0))
    db.session.add(tag)
    if screen is not None:
        screen.add_tag(tag)
    db.session.commit()
    return tag


# ---------------------------------------------------------------- general: stored in the database
def test_general_value_stored_in_database(client, login_as, db, dummy, dashboards):
    """A written value is a database row and survives a lost memory copy (server restart)."""
    login_as("admin")
    tag = make_tag(db, dummy, "Mode", "INT", screen=dashboards["public"])
    assert client.post(f"/api/datapoints/{tag.id}/write", json={"value": 42}).get_json()["value"] == 42
    row = DummyValue.query.filter_by(plc_id=dummy.id, slot=1).one()
    assert json.loads(row.value) == 42 and row.data_type == "INT"
    DummyConnection.reset()  # like a restart of the web server
    url = f"/api/dashboards/{dashboards['public'].id}/values"
    assert client.get(url).get_json()["values"][str(tag.id)]["value"] == 42


def test_general_value_same_for_every_client(app, db, dummy, dashboards, users):
    """Two different browsers (sessions) see the same value."""
    tag = make_tag(db, dummy, "Shared", "BOOL", screen=dashboards["public"])
    from tests.conftest import login  # noqa: PLC0415

    first, second = app.test_client(), app.test_client()

    login(first, "admin")
    login(second, "viewer")
    assert first.post(f"/api/datapoints/{tag.id}/write", json={"value": True}).status_code == 200
    url = f"/api/dashboards/{dashboards['public'].id}/values"
    assert second.get(url).get_json()["values"][str(tag.id)]["value"] is True


def test_general_value_survives_type_change(db, dummy):
    """After changing the data type the stored value is converted when it fits (else the default)."""
    tag = make_tag(db, dummy, "T", "UINT")
    db.session.add(DummyValue(plc_id=dummy.id, slot=1, data_type="INT", value="10"))
    db.session.commit()
    assert read_datapoints([tag])[0][tag.id]["value"] == 10       # INT 10 -> UINT 10
    real = make_tag(db, dummy, "R", "INT", slot=2)
    db.session.add(DummyValue(plc_id=dummy.id, slot=2, data_type="REAL", value="3.6"))
    neg = make_tag(db, dummy, "N", "UINT", slot=3)
    db.session.add(DummyValue(plc_id=dummy.id, slot=3, data_type="INT", value="-5"))
    text = make_tag(db, dummy, "S", "INT", slot=4)
    db.session.add(DummyValue(plc_id=dummy.id, slot=4, data_type="STRING", value='"abc"'))
    db.session.commit()
    values = read_datapoints([real, neg, text])[0]
    assert values[real.id]["value"] == 4          # rounded
    assert values[neg.id]["value"] == 0           # does not fit UINT -> default
    assert values[text.id]["value"] == 0          # not a number -> default


def test_bad_stored_value_falls_back_to_default(db, dummy):
    """Corrupt or out-of-range rows give the default value."""
    tag = make_tag(db, dummy, "T", "INT")
    db.session.add(DummyValue(plc_id=dummy.id, slot=1, data_type="INT", value="99999999"))
    db.session.commit()
    assert read_datapoints([tag])[0][tag.id]["value"] == 0


def test_dummy_values_without_app_context(app):
    """Outside an application context the memory copy is used."""
    conn = DummyConnection(99)
    addr = DummyAddress(start=1, data_type="INT")
    conn.write_tag(addr, 5)
    assert conn.read_tags({"a": addr})["a"].value == 5


def test_next_slot_forgets_old_values(db, dummy):
    """A new tag on a freed slot does not inherit the deleted tag's value."""
    tag = make_tag(db, dummy, "Old", "INT", slot=1)
    db.session.add(DummyValue(plc_id=dummy.id, slot=1, data_type="INT", value="7"))
    db.session.commit()
    db.session.delete(tag)
    db.session.commit()
    assert next_slot(dummy.id) == 1
    db.session.commit()
    assert DummyValue.query.count() == 0


def test_values_deleted_with_plc(db, dummy):
    """Stored values belong to their PLC."""
    db.session.add(DummyValue(plc_id=dummy.id, slot=1, data_type="INT", value="7"))
    db.session.commit()
    db.session.delete(dummy)
    db.session.commit()
    assert DummyValue.query.count() == 0


# ---------------------------------------------------------------- local: browser window only
def test_local_tag_not_read_or_written_by_server(client, login_as, db, dummy, dashboards):
    """The server sends the start value and the local spec, and refuses writes."""
    login_as("admin")
    tag = make_tag(db, dummy, "Show details", "BOOL", scope="local", screen=dashboards["public"],
                   true_text="Shown", false_text="Hidden")
    data = client.get(f"/api/dashboards/{dashboards['public'].id}/values").get_json()
    v = data["values"][str(tag.id)]
    assert v["quality"] == "local" and v["value"] is False and v["display"] == "Hidden"
    assert v["local"]["bool"] and v["local"]["true_text"] == "Shown" and v["local"]["initial"] is False
    assert str(dummy.id) not in data["plcs"]  # only local tags: no PLC badge
    resp = client.post(f"/api/datapoints/{tag.id}/write", json={"value": True})
    assert resp.status_code == 400 and "local" in resp.get_json()["error"]
    assert DummyValue.query.count() == 0


def test_local_spec(db, dummy):
    """The spec describes types, ranges and formats for hmi-local.js."""
    i = make_tag(db, dummy, "I", "INT", slot=1, scope="local", unit="rpm")
    s = make_tag(db, dummy, "S", "STRING", slot=2, scope="local")
    r = make_tag(db, dummy, "R", "REAL", slot=3, scope="local", decimals=1)
    assert i.local_spec()["range"] == [-32768, 32767] and i.local_spec()["integer"] and i.local_spec()["unit"] == "rpm"
    assert s.local_spec()["string_length"] == 10 and s.local_spec()["initial"] == ""
    assert r.local_spec()["range"] is None and r.local_spec()["decimals"] == 1
    assert i.to_dict()["dummy_scope"] == "local" and i.to_dict()["local"]["data_type"] == "INT"


def test_scope_only_for_dummy_plcs(db, plc, dummy):
    """Real PLC tags are never local, even with the column set."""
    real = DataPoint(plc=plc, label="Real", dummy_scope="local")
    real.set_address(DummyAddress(start=1, data_type="INT"))  # address type does not matter here
    tag = make_tag(db, dummy, "L", "INT", scope="local")
    assert tag.is_local and tag.scope_label == "local"
    assert not real.is_local and real.scope_label == ""


def test_tag_form_sets_scope(client, login_as, db, dummy, plc):
    """The tag form stores the kind for dummy tags; moving to a real PLC resets it."""
    login_as("admin")
    html = client.get("/tags/new").get_data(as_text=True)
    assert 'id="dummy-scope-row"' in html and "Local – own value in each browser window" in html
    client.post("/tags/new", data=tag_form(dummy, label="Panel", address="", data_type="BOOL", dummy_scope="local"))
    tag = DataPoint.query.filter_by(label="Panel").one()
    assert tag.is_local
    edit = client.get(f"/tags/{tag.id}/edit").get_data(as_text=True)
    assert re.search(r'<option selected value="local"|<option value="local" selected', edit)
    client.post(f"/tags/{tag.id}/edit", data=tag_form(plc, label="Panel", address="M1.0", data_type="BOOL",
                                                       dummy_scope="local"))
    db.session.refresh(tag)
    assert tag.dummy_scope == "general" and not tag.is_local


def test_tag_table_shows_kind(client, login_as, db, dummy):
    """The tag table marks general and local dummy tags."""
    login_as("admin")
    make_tag(db, dummy, "G", "INT", slot=1)
    make_tag(db, dummy, "L", "INT", slot=2, scope="local")
    html = client.get("/tags/").get_data(as_text=True)
    assert ">local</span>" in html and ">general</span>" in html and "hmi-local.js" in html
    values = client.get("/api/tags/values?ids=" + ",".join(str(t.id) for t in DataPoint.query)).get_json()["values"]
    assert {v["quality"] for v in values.values()} == {"good", "local"}


def test_designer_api_creates_local_tag(client, login_as, db, dummy):
    """The designer's New-tag dialog sends the kind."""
    login_as("admin")
    screen = Dashboard(name="Mimic", slug="mimic", kind="image")
    db.session.add(screen)
    db.session.commit()
    resp = client.post(f"/api/dashboards/{screen.id}/datapoints",
                       json={"label": "Toggle", "plc_id": dummy.id, "data_type": "BOOL", "writable": True,
                             "dummy_scope": "local"})
    assert resp.status_code == 201 and resp.get_json()["dummy_scope"] == "local"
    assert 'id="tag-scope-box"' in client.get(f"/screens/{screen.id}/designer").get_data(as_text=True)


def test_runtime_pages_load_local_script(client, login_as, db, dummy, dashboards):
    """Grid and image screens include hmi-local.js; grid tiles carry the spec, writable without permission."""
    tag = make_tag(db, dummy, "Mode", "INT", scope="local", screen=dashboards["public"])
    login_as("guest")  # Guest role: no write permission
    html = client.get("/screens/public").get_data(as_text=True)
    assert "hmi-local.js" in html and "data-local=" in html
    tile = re.search(r'<div class="hmi-tile[^>]*data-id="%d".*?</form>' % tag.id, html, re.S)
    assert tile and "hmi-write" in tile.group(0)
    image = Dashboard(name="Img", slug="img", kind="image", is_public=True)
    db.session.add(image)
    image.add_tag(tag)
    db.session.commit()
    page = client.get("/screens/img").get_data(as_text=True)
    assert "hmi-local.js" in page and '"dummy_scope": "local"' in page


# ---------------------------------------------------------------- alarms
def test_alarm_form_rejects_local_tag(client, login_as, db, dummy, dashboards):
    """Alarms cannot watch local tags."""
    login_as("admin")
    tag = make_tag(db, dummy, "Local", "INT", scope="local")
    html = client.post("/alarms/new", data={"name": "X", "kind": "tag", "datapoint_id": tag.id, "operator": "gt",
                                            "setpoint": "1", "severity": "low", "screen_id": 0, "plc_id": 0,
                                            "enabled": "y"}).get_data(as_text=True)
    assert "cannot raise alarms" in html and Alarm.query.count() == 0


def test_tag_with_alarm_cannot_become_local(client, login_as, db, dummy):
    """Making a tag local is refused while alarms use it."""
    login_as("admin")
    tag = make_tag(db, dummy, "Level", "INT")
    db.session.add(Alarm(name="High", datapoint=tag, operator="gt", setpoint="5"))
    db.session.commit()
    html = client.post(f"/tags/{tag.id}/edit", data=tag_form(dummy, label="Level", address="",
                                                              dummy_scope="local")).get_data(as_text=True)
    assert "cannot raise alarms" in html
    db.session.refresh(tag)
    assert not tag.is_local


def test_engine_skips_local_tags(db, dummy):
    """An alarm left on a local tag is never evaluated."""
    tag = make_tag(db, dummy, "Level", "INT")
    db.session.add(Alarm(name="Zero", datapoint=tag, operator="eq", setpoint="0"))
    db.session.commit()
    assert scan_alarms()["raised"] == 1
    tag.dummy_scope = "local"
    db.session.commit()
    stats = scan_alarms()
    assert stats["checked"] == 0


# ---------------------------------------------------------------- Excel
def test_excel_roundtrip_kind(client, login_as, db, dummy):
    """Export writes the kind; import sets it."""
    login_as("admin")
    make_tag(db, dummy, "G", "INT", slot=1)
    tag = make_tag(db, dummy, "L", "BOOL", slot=2, scope="local")
    wb = load_workbook(io.BytesIO(client.get("/tags/export").data))
    rows = list(wb["Tags"].iter_rows(values_only=True))
    col = rows[0].index("Dummy kind")
    assert {r[1]: r[col] for r in rows[1:]} == {"G": "general", "L": "local"}
    # flip the kinds and import
    ws = wb["Tags"]
    for row in ws.iter_rows(min_row=2):
        row[col].value = "local" if row[1].value == "G" else "general"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = client.post("/tags/import", data={"file": (buf, "tags.xlsx")}, content_type="multipart/form-data")
    assert resp.status_code in (200, 302)
    db.session.expire_all()
    assert DataPoint.query.filter_by(label="G").one().is_local
    assert not db.session.get(DataPoint, tag.id).is_local


def test_svg_import_keeps_kind(client, login_as, db, dummy):
    """Tags recreated by the SVG import keep their kind."""
    from app.dashboards.svg_import import ImportReport, _resolve_tag  # noqa: PLC0415

    screen = Dashboard(name="S", slug="s", kind="image")
    db.session.add(screen)
    db.session.commit()
    dp = _resolve_tag("77", {"label": "Win", "plc": "Sim", "data_type": "BOOL", "dummy_scope": "local"},
                      screen, ImportReport())
    assert dp is not None and dp.is_local


# ---------------------------------------------------------------- alarm scanning (water level scenario)
def test_water_level_alarms(client, login_as, db, dummy):
    """INT dummy tag with '> 90.0' and '≤ 10.0' alarms: written values raise and clear them."""
    from app.models import AlarmEvent  # noqa: PLC0415

    login_as("admin")
    tag = make_tag(db, dummy, "water_level", "INT")
    for name, op, sp in (("water is full", "gt", "90.0"), ("water is low", "le", "10.0")):
        db.session.add(Alarm(name=name, datapoint=tag, operator=op, setpoint=sp))
    db.session.commit()

    def active():
        scan_alarms()
        return sorted(e.name for e in AlarmEvent.query.filter(AlarmEvent.cleared_at.is_(None)))

    for value, expected in ((50, []), (95, ["water is full"]), (5, ["water is low"]), (10, ["water is low"]),
                            (11, [])):
        assert client.post(f"/api/datapoints/{tag.id}/write", json={"value": value}).status_code == 200
        assert active() == expected, value


def test_failing_plc_does_not_block_other_alarms(db, dummy, plc, monkeypatch):
    """An unexpected error while reading one PLC only affects that PLC's alarms."""
    from app.alarms import engine  # noqa: PLC0415
    from app.models import AlarmEvent  # noqa: PLC0415
    from app.plc import service  # noqa: PLC0415
    from tests.conftest import add_point  # noqa: PLC0415

    screen = Dashboard(name="S", slug="s")
    db.session.add(screen)
    db.session.commit()
    real = add_point(screen, plc, "Real", "DB1.DBW0:INT")
    tag = make_tag(db, dummy, "Level", "INT")
    db.session.add_all([Alarm(name="real", datapoint=real, operator="ge", setpoint="0"),
                        Alarm(name="dummy", datapoint=tag, operator="eq", setpoint="0")])
    db.session.commit()
    original = service.read_datapoints

    def broken(points):
        if any(p.plc_id == plc.id for p in points):
            raise RuntimeError("driver crashed")
        return original(points)

    monkeypatch.setattr(service, "read_datapoints", broken)
    stats = engine.scan_alarms()
    assert [e.name for e in AlarmEvent.query] == ["dummy"]
    assert stats["errors"] and "driver crashed" in stats["errors"][0]


def test_scanner_status_and_restart(app, client, login_as, monkeypatch):
    """The Alarms page shows the scanner health; a dead thread is restarted."""
    import threading  # noqa: PLC0415

    from app.alarms import engine  # noqa: PLC0415

    login_as("admin")
    engine.scan_alarms()
    html = client.get("/alarms/").get_data(as_text=True)
    assert 'id="scanner-status"' in html and "ALARM_SCANNER=0" in html  # testing config: thread off
    dead = threading.Thread(target=lambda: None)
    dead.start()
    dead.join()
    monkeypatch.setattr(engine.AlarmScanner, "thread", dead)
    monkeypatch.setattr(engine.AlarmScanner, "_run", staticmethod(lambda app, interval: None))
    assert engine.AlarmScanner.start(app) is True and engine.AlarmScanner.thread is not dead


def test_scan_if_due_never_blocks(app):
    """While a scan is running, the events API does not wait for it."""
    from app.alarms import engine  # noqa: PLC0415

    engine._last_scan["t"] = 0.0
    with engine._scan_lock:
        engine.scan_if_due()  # returns at once instead of dead-locking
