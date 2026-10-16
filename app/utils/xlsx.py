"""Excel (.xlsx) helpers shared by the tag and alarm export/import (openpyxl).

* :func:`build_workbook` writes one data sheet (bold, frozen header row with
  filters and sensible column widths) plus a *Help* sheet describing the
  columns, and returns the file as bytes.
* :func:`read_rows` reads the first sheet (or a sheet with a given name) of
  an uploaded workbook into ``(row number, {column: value})`` pairs using the
  header row, ignoring empty rows.
* :func:`text`, :func:`number`, :func:`boolean` convert cell values.
"""

from __future__ import annotations

import io
from datetime import datetime
from app.i18n.catalog import gettext

#: Largest accepted upload.
MAX_IMPORT_BYTES = 5 * 1024 * 1024
#: Most data rows accepted in one import.
MAX_IMPORT_ROWS = 5000

_TRUE = {"1", "true", "yes", "y", "on", "x", "✓", "enabled"}
_FALSE = {"0", "false", "no", "n", "off", "", "disabled"}


class ImportFileError(ValueError):
    """The uploaded file cannot be read as a workbook with the expected columns."""


def build_workbook(sheet_title: str, columns: list[tuple[str, int]], rows, help_rows: list[tuple[str, str]]) -> bytes:
    """Create an .xlsx file.

    Args:
        sheet_title: Name of the data sheet (e.g. ``Tags``).
        columns: ``(header, width)`` pairs.
        rows: Iterable of row value lists (same order as ``columns``).
        help_rows: ``(column, explanation)`` pairs for the *Help* sheet.

    Returns:
        The workbook as bytes.
    """
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    ws.append([name for name, _ in columns])
    for idx, (_, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    count = 0
    for row in rows:
        ws.append(list(row))
        count += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(count + 1, 2)}"

    info = wb.create_sheet("Help")
    info.append(["Column", "Meaning"])
    for cell in info[1]:
        cell.font, cell.fill = header_font, header_fill
    for name, meaning in help_rows:
        info.append([name, meaning])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    info.append([])
    info.append(["Exported", datetime.now().strftime("%Y-%m-%d %H:%M")])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def read_rows(stream, aliases: dict[str, str], required: tuple[str, ...], sheet_name: str | None = None):
    """Read an uploaded workbook into dictionaries keyed by canonical column names.

    Args:
        stream: File-like object (``request.files[...]``).
        aliases: ``lower-case header -> canonical name`` (include every canonical name itself).
        required: Canonical names that must be present in the header row.
        sheet_name: Preferred sheet (falls back to the first sheet).

    Returns:
        ``(columns, rows)``: the set of canonical columns found, and a list of
        ``(excel row number, {canonical name: value})`` for non-empty rows.

    Raises:
        ImportFileError: If the file is not a readable .xlsx, too large, has
            too many rows, or lacks required columns.
    """
    from openpyxl import load_workbook  # noqa: PLC0415

    data = stream.read(MAX_IMPORT_BYTES + 1)
    if not data:
        raise ImportFileError(gettext('The file is empty.'))
    if len(data) > MAX_IMPORT_BYTES:
        raise ImportFileError(gettext('The file is larger than {0} MB.', MAX_IMPORT_BYTES // (1024 * 1024)))
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - zipfile/openpyxl raise many types
        raise ImportFileError(gettext("This is not an Excel .xlsx file (save it as 'Excel Workbook (*.xlsx)').")) from exc
    try:
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            raise ImportFileError(gettext('The first sheet has no header row.'))
        mapping = {}
        for idx, name in enumerate(header):
            key = aliases.get(str(name or "").strip().lower())
            if key and key not in mapping.values():
                mapping[idx] = key
        missing = [name for name in required if name not in mapping.values()]
        if missing:
            raise ImportFileError(gettext('Missing column(s): {0}. Export a file first to get the right columns.', ', '.join(missing)))
        rows = []
        for number, values in enumerate(rows_iter, start=2):
            record = {key: (values[idx] if idx < len(values) else None) for idx, key in mapping.items()}
            if all(v is None or str(v).strip() == "" for v in record.values()):
                continue
            rows.append((number, record))
            if len(rows) > MAX_IMPORT_ROWS:
                raise ImportFileError(gettext('Too many rows (more than {0}). Split the file.', MAX_IMPORT_ROWS))
        return set(mapping.values()), rows
    finally:
        wb.close()


def text(value) -> str:
    """Cell value as stripped text (``""`` for empty cells).

    Args:
        value: Cell value.

    Returns:
        Text; whole floats such as ``3.0`` become ``"3"``.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    """Cell value as a number, ``None`` for an empty cell.

    Args:
        value: Cell value.

    Returns:
        ``int``/``float`` or ``None``.

    Raises:
        ValueError: If the text is not a number.
    """
    raw = text(value).replace(",", ".") if isinstance(value, str) else value
    if raw is None or raw == "":
        return None
    return float(raw) if not isinstance(raw, (int, float)) else raw


def boolean(value, default: bool = False) -> bool:
    """Cell value as a boolean (yes/no, true/false, 1/0, x).

    Args:
        value: Cell value.
        default: Result for an empty cell.

    Returns:
        ``True`` or ``False``.

    Raises:
        ValueError: For other text.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    if isinstance(value, bool):
        return value
    raw = text(value).lower()
    if raw in _TRUE:
        return True
    if raw in _FALSE:
        return False
    raise ValueError(f"'{value}' is not yes/no")


def yes_no(flag: bool) -> str:
    """``yes`` / ``no`` for export.

    Args:
        flag: Boolean.

    Returns:
        Text.
    """
    return "yes" if flag else "no"


#: MIME type of .xlsx files.
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def send_xlsx(data: bytes, name: str):
    """Send workbook bytes as a download named ``<name>-YYYYmmdd-HHMM.xlsx``.

    Args:
        data: Workbook bytes.
        name: File name prefix.

    Returns:
        A Flask response.
    """
    from flask import send_file  # noqa: PLC0415

    filename = f"{name}-{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    return send_file(io.BytesIO(data), mimetype=XLSX_MIME, as_attachment=True, download_name=filename)

