"""Tests for the tag history: sampling, the History page, Excel export and trend charts."""

import io
import json
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app.history import logger as hl
from app.history.data import parse_tag_ids, series_points
from app.models import PLC, DataPoint, Role, TagSample, TrendChart
from app.models.history import TagSample as Sample
from app.plc.dummy import DummyAddress, DummyConnection
from app.plc.service import write_datapoint
from app.utils.timezones import format_local, parse_local, zone_by_name
from tests.conftest import add_point
from tests.test_tag_features import tag_form

T0 = 1_800_000_000  # a fixed epoch second (2027-01-15 08:00:00 UTC), multiple of 10


@pytest.fixture(autouse=True)
def _reset():
    """Start each test with an empty logger schedule and dummy memory."""
    hl.reset_schedule()
    DummyConnection.reset()
    yield
    hl.reset_schedule()
    DummyConnection.reset()


@pytest.fixture()
def dummy(db):
    """A dummy PLC."""
    plc = PLC(name="Sim", ip_address="", protocol="dummy")
    db.session.add(plc)
    db.session.commit()
    return plc


def dummy_tag(db, plc, label, slot=1, dtype="INT", scope="general", interval=10, logged=True, **kw):
    """Create a dummy tag with history logging.

    Args:
        db: Database fixture.
        plc: Dummy PLC.
        label: Label.
        slot: Dummy slot.
        dtype: Data type.
        scope: general / local.
        interval: Sample interval (s).
        logged: History logging on.
        **kw: More DataPoint columns.

    Returns:
        The tag.
    """
    tag = DataPoint(plc=plc, label=label, writable=True, dummy_scope=scope, log_enabled=logged, log_interval=interval,
                    **kw)
    tag.set_address(DummyAddress(start=slot, data_type=dtype, string_length=10 if dtype == "STRING" else 0))
    db.session.add(tag)
    db.session.commit()
    return tag


def add_samples(db, tag, values, start=T0, step=10):
    """Insert samples directly.

    Args:
        db: Database fixture.
        tag: The tag.
        values: Values in time order.
        start: Epoch seconds of the first sample.
        step: Seconds between samples.
    """
    for i, v in enumerate(values):
        db.session.add(Sample.from_value(tag.id, datetime.fromtimestamp(start + i * step, tz=timezone.utc), v))
    db.session.commit()


# ---------------------------------------------------------------- sampling
def test_sampling_respects_interval_and_alignment(db, dummy):
    """Tags are sampled when due; the next sample is aligned to the interval."""
    fast = dummy_tag(db, dummy, "fast", slot=1, interval=10)
    slow = dummy_tag(db, dummy, "slow", slot=2, interval=60)
    dummy_tag(db, dummy, "off", slot=3, logged=False)
    write_datapoint(fast, 42)
    assert hl.sample_due_tags(T0 + 3)["saved"] == 2      # first run: every logged tag
    assert hl.sample_due_tags(T0 + 5)["saved"] == 0      # nothing due yet
    stats = hl.sample_due_tags(T0 + 10)                  # aligned boundary of the 10 s tag
    assert stats["due"] == 1 and stats["saved"] == 1
    assert hl.sample_due_tags(T0 + 60)["saved"] == 2     # both due at the minute
    samples = TagSample.query.filter_by(datapoint_id=fast.id).order_by(TagSample.ts).all()
    assert [s.value for s in samples] == [42.0, 42.0, 42.0]
    assert samples[1].ts.replace(tzinfo=timezone.utc) == datetime.fromtimestamp(T0 + 10, tz=timezone.utc)
    assert TagSample.query.filter_by(datapoint_id=slow.id).count() == 2


def test_sampling_value_kinds_and_local_tags(db, dummy):
    """BOOL is stored as 0/1, STRING as text; local dummy tags are never logged."""
    b = dummy_tag(db, dummy, "pump", slot=1, dtype="BOOL")
    s = dummy_tag(db, dummy, "msg", slot=2, dtype="STRING")
    local = dummy_tag(db, dummy, "mine", slot=3, scope="local")
    write_datapoint(b, True)
    write_datapoint(s, "hello")
    hl.sample_due_tags(T0)
    assert TagSample.query.filter_by(datapoint_id=b.id).one().value_num == 1.0
    sample = TagSample.query.filter_by(datapoint_id=s.id).one()
    assert sample.value_text == "hello" and sample.display() == "hello"
    assert TagSample.query.filter_by(datapoint_id=local.id).count() == 0
    assert TagSample.query.filter_by(datapoint_id=b.id).one().display() == b.true_text


def test_sampling_skips_bad_values_and_isolates_plcs(db, dummy, plc, monkeypatch):
    """A PLC whose read fails does not stop the samples of other PLCs."""
    good = dummy_tag(db, dummy, "good", slot=1)
    broken = add_point(None, plc, "broken", "DB1.DBW0", log_enabled=True, log_interval=10)
    from app.plc import service

    real = service.read_datapoints

    def fake(points):
        """Fail for the real PLC."""
        if points[0].plc_id == plc.id:
            raise RuntimeError("cable cut")
        return real(points)

    monkeypatch.setattr(service, "read_datapoints", fake)
    stats = hl.sample_due_tags(T0)
    assert stats["saved"] == 1 and stats["skipped"] == 1 and "cable cut" in stats["errors"][0]
    assert TagSample.query.filter_by(datapoint_id=good.id).count() == 1
    assert TagSample.query.filter_by(datapoint_id=broken.id).count() == 0


def test_interval_change_resamples(db, dummy):
    """Changing a tag's interval takes effect immediately."""
    tag = dummy_tag(db, dummy, "t", interval=60)
    hl.sample_due_tags(T0)
    tag.log_interval = 5
    db.session.commit()
    assert hl.sample_due_tags(T0 + 1)["saved"] == 1
    assert hl.sample_due_tags(T0 + 5)["saved"] == 1


def test_purge_old_samples(db, dummy, app):
    """Samples older than the retention period are deleted; 0 keeps everything."""
    tag = dummy_tag(db, dummy, "t")
    now = datetime.now(timezone.utc).timestamp()
    add_samples(db, tag, [1, 2], start=int(now - 40 * 86400), step=86400 * 30)  # 40 and 10 days old
    assert hl.purge_old_samples(0) == 0
    assert hl.purge_old_samples(30) == 1
    assert TagSample.query.count() == 1
    from app.models.setting import set_setting

    set_setting(hl.RETENTION_SETTING, 5)
    db.session.commit()
    assert hl.retention_days() == 5 and hl.purge_old_samples() == 1


def test_logger_thread_starts(app, db):
    """The thread starts on demand and reports its state."""
    hl.HistoryLogger.thread = None
    with app.app_context():
        assert hl.logger_status()["thread"] == "off"  # HISTORY_LOGGER is False in tests
    app.config["HISTORY_LOGGER"] = True
    app.config["HISTORY_TICK"] = 0.2
    try:
        assert hl.HistoryLogger.start(app) is True
        assert hl.HistoryLogger.start(app) is False
        assert hl.logger_status()["thread"] == "running"
    finally:
        hl.HistoryLogger.stop()
        app.config["HISTORY_LOGGER"] = False
    assert hl.HistoryLogger.thread is None


def test_deleting_tag_removes_samples_and_chart_series(db, dummy):
    """Samples and chart series of a deleted tag disappear (also on SQLite)."""
    a = dummy_tag(db, dummy, "a", slot=1)
    b = dummy_tag(db, dummy, "b", slot=2)
    add_samples(db, a, [1, 2, 3])
    chart = TrendChart(name="c")
    chart.series = [{"tag_id": a.id, "color": "#ff0000"}, {"tag_id": b.id, "color": "#00ff00"}]
    db.session.add(chart)
    db.session.commit()
    db.session.delete(a)
    db.session.commit()
    assert TagSample.query.count() == 0
    db.session.refresh(chart)
    assert [s["tag_id"] for s in chart.series] == [b.id]


# ---------------------------------------------------------------- helpers
def test_parse_helpers():
    """Tag id lists and local dates are parsed safely."""
    assert parse_tag_ids("3, 1,x,3", ["7"]) == [3, 1, 7]
    zone = zone_by_name("Asia/Jerusalem")
    assert zone is not None and zone_by_name("Not/AZone") is None
    utc = parse_local("2026-07-01T12:00", zone)
    assert utc == datetime(2026, 7, 1, 9, 0, tzinfo=timezone.utc)  # IDT = UTC+3
    assert parse_local("garbage", zone) is None
    assert format_local(utc, zone) == "2026-07-01 12:00:00"


def test_series_points_gaps_and_reduction(db, dummy):
    """Gaps break the line; large ranges are reduced to min/max per slot."""
    tag = dummy_tag(db, dummy, "t", interval=10)
    add_samples(db, tag, [1, 2, 3])
    add_samples(db, tag, [4, 5], start=T0 + 1000)
    start = datetime.fromtimestamp(T0 - 10, tz=timezone.utc)
    end = datetime.fromtimestamp(T0 + 2000, tz=timezone.utc)
    res = series_points(tag, start, end)
    values = [p[1] for p in res["points"]]
    assert values == [1, 2, 3, None, 4, 5] and not res["reduced"]

    many = [((i * 7) % 50) + (100 if i == 333 else 0) for i in range(1000)]
    add_samples(db, tag, many, start=T0 + 5000, step=10)
    res = series_points(tag, datetime.fromtimestamp(T0 + 5000, tz=timezone.utc),
                        datetime.fromtimestamp(T0 + 15000, tz=timezone.utc), max_points=100)
    assert res["reduced"] and res["count"] == 1000 and len(res["points"]) <= 110
    assert max(p[1] for p in res["points"] if p[1] is not None) == max(many)  # the peak survives


# ---------------------------------------------------------------- tag form / Excel columns
def test_tag_form_history_settings(client, login_as, plc, dummy, db):
    """The tag form saves the history settings and refuses local dummy tags."""
    login_as("admin")
    resp = client.post("/tags/new", data=tag_form(plc, label="Flow", log_enabled="y", log_interval=15))
    assert resp.status_code == 302
    dp = DataPoint.query.filter_by(label="Flow").one()
    assert dp.log_enabled and dp.log_interval == 15
    html = client.get(f"/tags/{dp.id}/edit").get_data(as_text=True)
    assert 'id="history-settings"' in html and 'name="log_interval"' in html
    html = client.post("/tags/new", data=tag_form(dummy, label="L", address="", dummy_scope="local",
                                                    log_enabled="y", log_interval=5)).get_data(as_text=True)
    assert "its history cannot be saved" in html
    html = client.post("/tags/new", data=tag_form(plc, label="X", log_enabled="y", log_interval=0)).get_data(
        as_text=True)
    assert DataPoint.query.filter_by(label="X").first() is None and "between 1 and 86400" in html
    assert "tag-logged" in client.get("/tags/").get_data(as_text=True)


def test_tag_excel_history_columns(client, login_as, plc, db):
    """The tag Excel export/import carries the history columns."""
    login_as("admin")
    dp = add_point(None, plc, "Level", "DB1.DBW2", log_enabled=True, log_interval=30)
    wb = load_workbook(io.BytesIO(client.get("/tags/export").data))
    ws = wb["Tags"]
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert row[header.index("History")] == "yes" and row[header.index("History interval (s)")] == 30
    ws.cell(row=2, column=header.index("History interval (s)") + 1, value="5")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    client.post("/tags/import", data={"file": (buf, "t.xlsx")}, content_type="multipart/form-data")
    db.session.refresh(dp)
    assert dp.log_interval == 5 and dp.log_enabled


# ---------------------------------------------------------------- pages and permissions
def test_permissions(client, login_as, users, db):
    """Viewers see the history; guests do not; only screen managers edit charts."""
    assert Role.query.filter_by(name="Viewer").one().has_permission("view_history")
    login_as("guest")
    assert client.get("/history/").status_code == 403
    assert client.get("/history/charts").status_code == 403
    client.post("/auth/logout")
    login_as("viewer")
    assert client.get("/history/").status_code == 200
    html = client.get("/history/charts").get_data(as_text=True)
    assert 'id="new-chart"' not in html
    assert client.get("/history/charts/new").status_code == 403
    assert 'id="nav-history"' in client.get("/screens/").get_data(as_text=True)


def test_history_page_filters_and_time_zone(client, login_as, db, dummy):
    """The table filters by tag and local time range and shows full local dates."""
    a = dummy_tag(db, dummy, "alpha", slot=1)
    b = dummy_tag(db, dummy, "beta", slot=2)
    add_samples(db, a, [1.5, 2.5, 3.5])
    add_samples(db, b, [7, 8])
    login_as("viewer")
    client.set_cookie("hmi_tz", "Asia%2FJerusalem")  # as written by encodeURIComponent()
    html = client.get(f"/history/?tags={a.id}").get_data(as_text=True)
    assert "Asia/Jerusalem" in html and "alpha" in html and "beta</td>" not in html
    assert "2027-01-15 10:00:20" in html  # 08:00:20 UTC shown in Israel winter time (UTC+2)
    html = client.get(f"/history/?t={a.id}&t={b.id}&from=2027-01-15T10:00:10&to=2027-01-15T10:00:10").get_data(
        as_text=True)
    body = html.split('id="samples-table"')[1].split("</tbody>")[0]
    assert body.count("<tr>") - 1 == 2, body  # 10:00:10 of alpha and beta
    assert "1–2 of 2" in html


def test_excel_export_layouts(client, login_as, db, dummy):
    """Both export layouts contain the full local date and the values."""
    a = dummy_tag(db, dummy, "alpha", slot=1, unit="°C")
    b = dummy_tag(db, dummy, "beta", slot=2, dtype="BOOL")
    add_samples(db, a, [1.5, 2.5])
    add_samples(db, b, [True, False])
    login_as("viewer")
    client.set_cookie("hmi_tz", "UTC")
    resp = client.get(f"/history/export.xlsx?tags={a.id},{b.id}")
    assert resp.headers["Content-Disposition"].startswith("attachment; filename=tag-history-")
    ws = load_workbook(io.BytesIO(resp.data))["History"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("Date & time", "Tag", "PLC", "Value")
    assert rows[1][0] == datetime(2027, 1, 15, 8, 0, 0) and rows[1][1] == "alpha" and rows[1][3] == 1.5
    assert len(rows) == 5
    wide = load_workbook(io.BytesIO(client.get(f"/history/export.xlsx?tags={a.id},{b.id}&layout=wide").data))
    rows = list(wide["History"].iter_rows(values_only=True))
    assert rows[0] == ("Date & time", "alpha [°C]", "beta")
    assert rows[1] == (datetime(2027, 1, 15, 8, 0, 0), 1.5, True)
    assert rows[2][1:] == (2.5, False)
    info = dict(wide["Info"].iter_rows(values_only=True))
    assert info["Time zone"] == "UTC" and info["Rows"] == 2


def test_series_api(client, login_as, db, dummy):
    """The chart API returns epoch-ms points per tag."""
    a = dummy_tag(db, dummy, "alpha", slot=1)
    add_samples(db, a, [1, 2])
    login_as("viewer")
    data = client.get(f"/history/api/series?tags={a.id}&start={(T0 - 5) * 1000}&end={(T0 + 60) * 1000}").get_json()
    assert data["series"][str(a.id)]["points"] == [[T0 * 1000, 1.0], [(T0 + 10) * 1000, 2.0]]
    bad = client.get(f"/history/api/series?tags={a.id}&start={T0 * 1000}&end={T0 * 1000}")
    assert bad.status_code == 400
    client.post("/auth/logout")
    assert client.get(f"/history/api/series?tags={a.id}").status_code in (302, 401)


# ---------------------------------------------------------------- charts
def test_chart_crud(client, login_as, db, dummy):
    """Charts are created with coloured series of logged tags, edited, viewed and deleted."""
    a = dummy_tag(db, dummy, "alpha", slot=1)
    b = dummy_tag(db, dummy, "beta", slot=2)
    s = dummy_tag(db, dummy, "text", slot=3, dtype="STRING")
    off = dummy_tag(db, dummy, "off", slot=4, logged=False)
    login_as("admin")
    page = client.get(f"/history/charts/new?tags={a.id},{b.id}&colors=%23112233,bad").get_data(as_text=True)
    assert "#112233" in page and "alpha" in page and '"label": "text"' not in page

    def post(series, name="Boiler"):
        """Submit the chart form."""
        return client.post("/history/charts/new", data={"name": name, "range_key": "8h",
                                                        "series": json.dumps(series)})

    for bad, msg in (([], "Add at least one tag."),
                     ([{"tag_id": off.id, "color": "#000000"}], "has no history logging"),
                     ([{"tag_id": s.id, "color": "#000000"}], "has no history logging"),
                     ([{"tag_id": a.id, "color": "red"}], "Invalid colour"),
                     ([{"tag_id": a.id, "color": "#000000"}, {"tag_id": a.id, "color": "#ffffff"}], "twice")):
        assert msg in post(bad).get_data(as_text=True)
    resp = post([{"tag_id": a.id, "color": "#FF0000", "axis": "left"}, {"tag_id": b.id, "color": "#00ff00", "axis": "right"}])
    chart = TrendChart.query.one()
    assert resp.headers["Location"].endswith(f"/history/charts/{chart.id}")
    assert chart.series == [{"tag_id": a.id, "color": "#ff0000", "axis": "left"},
                            {"tag_id": b.id, "color": "#00ff00", "axis": "right"}]
    assert "already exists" in post([{"tag_id": a.id, "color": "#000000"}]).get_data(as_text=True)
    html = client.get(f"/history/charts/{chart.id}").get_data(as_text=True)
    assert "chart.umd.min.js" in html and 'id="save-image"' in html and '"color": "#00ff00"' in html
    client.post(f"/history/charts/{chart.id}/edit", data={"name": "Boiler 2", "range_key": "1h",
                                                         "series": json.dumps([{"tag_id": b.id, "color": "#0000ff"}])})
    db.session.refresh(chart)
    assert chart.name == "Boiler 2" and chart.series[0]["color"] == "#0000ff"
    assert "Boiler 2" in client.get("/history/charts").get_data(as_text=True)
    client.post(f"/history/charts/{chart.id}/delete")
    assert TrendChart.query.count() == 0


def test_quick_chart(client, login_as, db, dummy):
    """Viewers can draw logged tags without saving; unusable tags are refused."""
    a = dummy_tag(db, dummy, "alpha", slot=1)
    off = dummy_tag(db, dummy, "off", slot=2, logged=False)
    login_as("viewer")
    html = client.get(f"/history/charts/quick?t={a.id}&range=24h").get_data(as_text=True)
    assert "Quick chart" in html and '"range": "24h"' in html and 'id="save-quick"' not in html
    assert client.get(f"/history/charts/quick?tags={off.id}").status_code == 302


def test_static_chart_files(client):
    """Chart.js and the date adapter are served."""
    for path in ("/static/vendor/chartjs/chart.umd.min.js", "/static/vendor/chartjs/chartjs-adapter-date-fns.bundle.min.js",
                 "/static/js/trend_chart.js", "/static/js/chart_form.js", "/static/js/history.js"):
        assert client.get(path).status_code == 200


def test_settings_retention(client, login_as, db):
    """Administrators set the retention period in System settings."""
    login_as("admin")
    html = client.get("/admin/settings").get_data(as_text=True)
    assert 'name="history_retention_days"' in html and 'value="90"' in html
    client.post("/admin/settings", data={"public_home_screen_id": 0, "history_retention_days": 14})
    assert hl.retention_days() == 14
    assert hl.purge_old_samples() == 0
    assert datetime.now(timezone.utc) - timedelta(days=14)  # sanity
