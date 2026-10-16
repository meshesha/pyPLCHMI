"""Queries on the tag history: filtered sample lists, chart series and Excel export.

* :func:`parse_tag_ids` reads the selected tags of a filter (``tags=1,4,7``).
* :func:`samples_query` builds the filtered query of the History page.
* :func:`series_points` returns the points of one chart series, reduced to
  at most ``max_points`` (minimum and maximum of each time slot, so peaks
  stay visible) and with gaps where the logger had no value.
* :func:`export_workbook` writes the filtered samples to an .xlsx file, as a
  list (one row per sample) or as a table with one column per tag.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

from app.i18n.catalog import gettext
from app.models import DataPoint, TagSample
from app.utils.timezones import to_local, zone_label

#: Most rows written to one Excel file (Excel's limit is 1,048,576).
MAX_EXPORT_ROWS = 1_000_000
#: Default and largest number of points per chart series.
DEFAULT_POINTS, MAX_POINTS = 2000, 10000
#: Most tags in one history filter / chart request.
MAX_TAGS = 50


def parse_tag_ids(*values) -> list[int]:
    """Collect tag ids from ``"1,2,3"`` strings and lists of strings.

    Args:
        *values: Strings or lists (``request.args.get("tags")``, ``getlist("t")``).

    Returns:
        Unique ids in the given order (at most :data:`MAX_TAGS`).
    """
    ids: list[int] = []
    for value in values:
        items = value if isinstance(value, (list, tuple)) else [value]
        for item in items:
            for part in str(item or "").split(","):
                part = part.strip()
                if part.isdigit() and int(part) not in ids:
                    ids.append(int(part))
    return ids[:MAX_TAGS]


def history_tags() -> list[DataPoint]:
    """Tags that can be chosen on the History page: logged now, or with saved samples.

    Returns:
        Tags ordered by label.
    """
    from app.extensions import db  # noqa: PLC0415

    with_samples = {row[0] for row in db.session.query(TagSample.datapoint_id).distinct()}
    tags = DataPoint.query.filter(
        (DataPoint.log_enabled.is_(True)) | (DataPoint.id.in_(with_samples or [0]))
    ).order_by(DataPoint.label).all()
    return [dp for dp in tags if dp.is_logged or dp.id in with_samples]


def chart_tags() -> list[DataPoint]:
    """Tags that can be put on a chart: history logging enabled and numeric / BOOL.

    Returns:
        Tags ordered by label.
    """
    return [dp for dp in DataPoint.query.filter_by(log_enabled=True).order_by(DataPoint.label).all()
            if dp.is_logged and dp.is_chartable]


def samples_query(tag_ids: list[int], start: datetime | None, end: datetime | None, newest_first: bool = True):
    """Samples of the given tags within a time range.

    Args:
        tag_ids: Tag ids (empty = all tags).
        start: First time (UTC, inclusive) or ``None``.
        end: Last time (UTC, inclusive) or ``None``.
        newest_first: Sort order of the time column.

    Returns:
        A SQLAlchemy query of :class:`TagSample`.
    """
    query = TagSample.query
    if tag_ids:
        query = query.filter(TagSample.datapoint_id.in_(tag_ids))
    if start is not None:
        query = query.filter(TagSample.ts >= start)
    if end is not None:
        query = query.filter(TagSample.ts <= end)
    order = TagSample.ts.desc() if newest_first else TagSample.ts.asc()
    return query.order_by(order, TagSample.datapoint_id.asc())


def _ms(ts: datetime) -> int:
    """Epoch milliseconds of a datetime (naive values count as UTC).

    Args:
        ts: The datetime.

    Returns:
        Milliseconds since 1970-01-01 UTC.
    """
    if ts.tzinfo is None:
        ts = ts.replace(tzinfo=timezone.utc)
    return int(ts.timestamp() * 1000)


def series_points(dp: DataPoint, start: datetime, end: datetime, max_points: int = DEFAULT_POINTS) -> dict:
    """Points of one tag for a chart.

    Args:
        dp: The tag.
        start: Range start (UTC).
        end: Range end (UTC).
        max_points: Most points returned (larger ranges are reduced).

    Returns:
        ``{"points": [[ms, value or None], …], "count": samples in range, "reduced": bool}``.
        ``None`` values break the line where samples are missing (a pause longer
        than three sample intervals and three times the previous spacing).
    """
    from app.extensions import db  # noqa: PLC0415

    max_points = max(10, min(MAX_POINTS, int(max_points)))
    base = db.session.query(TagSample.ts, TagSample.value_num).filter(
        TagSample.datapoint_id == dp.id, TagSample.ts >= start, TagSample.ts <= end,
        TagSample.value_num.isnot(None)).order_by(TagSample.ts.asc())
    count = base.count()
    gap_ms = max(3 * int(dp.log_interval or 60), 3) * 1000 + 1000
    points: list = []
    if count <= max_points:
        last, step = None, None
        for ts, value in base.yield_per(5000):
            ms = _ms(ts)
            if last is not None:
                delta = ms - last
                if delta > max(gap_ms, 3 * step if step else 60_000):
                    points.append([last + 1, None])
                    step = None
                else:
                    step = delta
            points.append([ms, value])
            last = ms
        return {"points": points, "count": count, "reduced": False}

    # Reduce: per time slot keep the minimum and the maximum (in time order).
    slots = max(1, max_points // 2)
    t0, t1 = _ms(start), _ms(end)
    width = max(1.0, (t1 - t0) / slots)
    gap_ms = max(gap_ms, int(width * 2))
    current, lo, hi, last = None, None, None, None

    def flush():
        """Append the min/max points of the finished slot."""
        nonlocal last
        if lo is None:
            return
        first, second = (lo, hi) if lo[0] <= hi[0] else (hi, lo)
        if last is not None and first[0] - last > gap_ms:
            points.append([last + 1, None])
        points.append(list(first))
        if second is not first and second[0] != first[0]:
            points.append(list(second))
        last = points[-1][0]

    for ts, value in base.yield_per(10000):
        ms = _ms(ts)
        slot = int((ms - t0) // width)
        if slot != current:
            flush()
            current, lo, hi = slot, (ms, value), (ms, value)
            continue
        if value < lo[1]:
            lo = (ms, value)
        if value > hi[1]:
            hi = (ms, value)
    flush()
    return {"points": points, "count": count, "reduced": True}


def _header(dp: DataPoint) -> str:
    """Column title of a tag in the one-column-per-tag layout.

    Args:
        dp: The tag.

    Returns:
        ``Label [unit]``.
    """
    return f"{dp.label} [{dp.unit}]" if dp.unit else dp.label


def _cell_value(sample_text, sample_num, dp: DataPoint):
    """Excel value of a sample: number, TRUE/FALSE for BOOL tags, or text.

    Args:
        sample_text: ``value_text`` of the sample.
        sample_num: ``value_num`` of the sample.
        dp: The tag.

    Returns:
        A value openpyxl can write.
    """
    if sample_text is not None:
        return sample_text
    if sample_num is not None and dp is not None and dp.is_bool:
        return bool(sample_num)
    return sample_num


def export_workbook(tag_ids: list[int], start: datetime | None, end: datetime | None, zone,
                    layout: str = "list") -> tuple[bytes, int, bool]:
    """Write samples to an .xlsx workbook (streamed, so large histories fit in memory).

    Args:
        tag_ids: Tags to export (empty = all).
        start: Range start (UTC) or ``None``.
        end: Range end (UTC) or ``None``.
        zone: Time zone of the *Date & time* column.
        layout: ``"list"`` – one row per sample (Date & time, Tag, PLC, Value,
            Unit, Address); ``"wide"`` – one row per time and one column per tag.

    Returns:
        ``(workbook bytes, data rows written, truncated)``.
    """
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.cell import WriteOnlyCell  # noqa: PLC0415
    from openpyxl.styles import Font, PatternFill  # noqa: PLC0415

    from app.extensions import db  # noqa: PLC0415

    tags = {dp.id: dp for dp in (DataPoint.query.filter(DataPoint.id.in_(tag_ids)).all() if tag_ids
                                 else DataPoint.query.all())}
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(gettext("History"))
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")

    def header(titles):
        """Write a styled header row."""
        cells = []
        for title in titles:
            cell = WriteOnlyCell(ws, value=title)
            cell.font, cell.fill = bold, fill
            cells.append(cell)
        ws.append(cells)

    query = db.session.query(TagSample.ts, TagSample.datapoint_id, TagSample.value_num, TagSample.value_text)
    if tag_ids:
        query = query.filter(TagSample.datapoint_id.in_(tag_ids))
    if start is not None:
        query = query.filter(TagSample.ts >= start)
    if end is not None:
        query = query.filter(TagSample.ts <= end)
    query = query.order_by(TagSample.ts.asc(), TagSample.datapoint_id.asc())

    rows, truncated = 0, False
    if layout == "wide":
        ordered = [tags[i] for i in (tag_ids or sorted(tags, key=lambda i: tags[i].label.lower())) if i in tags]
        position = {dp.id: idx for idx, dp in enumerate(ordered)}
        ws.column_dimensions["A"].width = 21
        for idx in range(len(ordered)):
            ws.column_dimensions[_col(idx + 2)].width = 16
        ws.freeze_panes = "B2"
        header([gettext("Date & time")] + [_header(dp) for dp in ordered])
        current_ts, row = None, None
        for ts, tag_id, num, text in query.yield_per(10000):
            if tag_id not in position:
                continue
            if ts != current_ts:
                if row is not None:
                    ws.append(row)
                    rows += 1
                    if rows >= MAX_EXPORT_ROWS:
                        truncated, row = True, None
                        break
                current_ts = ts
                row = [to_local(ts, zone).replace(tzinfo=None)] + [None] * len(ordered)
            row[1 + position[tag_id]] = _cell_value(text, num, tags[tag_id])
        if row is not None:
            ws.append(row)
            rows += 1
    else:
        for letter, width in zip("ABCDEF", (21, 28, 18, 14, 10, 22)):
            ws.column_dimensions[letter].width = width
        ws.freeze_panes = "A2"
        header([gettext("Date & time"), gettext("Tag"), gettext("PLC"), gettext("Value"), gettext("Unit"),
                gettext("Address")])
        for ts, tag_id, num, text in query.yield_per(10000):
            dp = tags.get(tag_id)
            if dp is None:
                continue
            ws.append([to_local(ts, zone).replace(tzinfo=None), dp.label, dp.plc.name if dp.plc else "",
                       _cell_value(text, num, dp), dp.unit, dp.address])
            rows += 1
            if rows >= MAX_EXPORT_ROWS:
                truncated = True
                break

    info = wb.create_sheet(gettext("Info"))
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70
    info.append([gettext("Exported"), datetime.now(zone).strftime("%Y-%m-%d %H:%M:%S")])
    info.append([gettext("Time zone"), zone_label(zone)])
    info.append([gettext("From"), to_local(start, zone).strftime("%Y-%m-%d %H:%M:%S") if start else gettext("first sample")])
    info.append([gettext("To"), to_local(end, zone).strftime("%Y-%m-%d %H:%M:%S") if end else gettext("last sample")])
    info.append([gettext("Tags"), ", ".join(dp.label for dp in tags.values()) if tag_ids else gettext("all")])
    info.append([gettext("Rows"), rows])
    if truncated:
        info.append([gettext("Note"), gettext("Only the first {0} rows were exported – choose a shorter time range.",
                                              MAX_EXPORT_ROWS)])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), rows, truncated


def _col(index: int) -> str:
    """Excel column letter of a 1-based column number.

    Args:
        index: Column number.

    Returns:
        ``A``, ``B`` … ``AA``.
    """
    from openpyxl.utils import get_column_letter  # noqa: PLC0415

    return get_column_letter(index)


def default_range(seconds: int) -> tuple[datetime, datetime]:
    """The last ``seconds`` up to now.

    Args:
        seconds: Length of the range.

    Returns:
        ``(start, end)`` in UTC.
    """
    end = datetime.now(timezone.utc)
    return end - timedelta(seconds=seconds), end


def tags_with_history(tag_ids) -> set[int]:
    """Which of the given tags have history: logging enabled or samples saved.

    Args:
        tag_ids: Tag ids (e.g. the tags of one screen).

    Returns:
        Set of ids.
    """
    from app.extensions import db  # noqa: PLC0415

    ids = [int(i) for i in tag_ids]
    if not ids:
        return set()
    logged = {dp.id for dp in DataPoint.query.filter(DataPoint.id.in_(ids), DataPoint.log_enabled.is_(True)).all()
              if not dp.is_local}
    sampled = {row[0] for row in db.session.query(TagSample.datapoint_id).filter(
        TagSample.datapoint_id.in_(ids)).distinct()}
    return logged | sampled
