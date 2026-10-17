"""Tests for the built-in Operator role and the per-tag write access (Everyone / selected roles)."""

import io
import re

from openpyxl import Workbook, load_workbook

from app.models import DataPoint, Role
from app.models.user import DEFAULT_ROLES, Perm
from tests.conftest import login, make_user

BASE = {"plc_id": None, "data_type": "REAL", "widget": "value", "width": 1, "scale": 1, "offset": 0,
        "decimals": 2, "true_text": "ON", "false_text": "OFF"}


def role(name: str) -> Role:
    """The role with this name."""
    return Role.query.filter_by(name=name).one()


def limit_to(dp: DataPoint, *names: str) -> None:
    """Limit writing of a tag to the named roles (and commit)."""
    from app.extensions import db  # noqa: PLC0415

    dp.write_access = "roles"
    dp.write_roles = [role(n) for n in names]
    db.session.commit()


def test_operator_is_a_builtin_role(app, client, login_as, users):
    """Operator exists after seeding, can write values, and cannot be deleted."""
    op = role("Operator")
    assert "Operator" in DEFAULT_ROLES
    assert op.permission_codes == {Perm.VIEW_SCREENS, Perm.VIEW_PUBLIC_SCREENS, Perm.VIEW_HISTORY,
                                   Perm.WRITE_VALUES}
    login_as("admin")
    client.post(f"/admin/roles/{op.id}/delete")
    assert Role.query.filter_by(name="Operator").count() == 1
    assert "Operator" in client.get("/admin/roles").get_data(as_text=True)


def test_may_write_rules(app, dashboards, users):
    """Everyone vs. selected roles; read-only tags; local dummy tags; visitors."""
    from app.models.user import AnonymousUser  # noqa: PLC0415

    operator = make_user("op", "Operator")
    sp = dashboards["points"]["setpoint"]
    temp = dashboards["points"]["temp"]
    assert sp.may_write(operator) and sp.may_write(users["admin"]) and not sp.may_write(users["viewer"])
    assert not temp.may_write(users["admin"])  # not writable
    limit_to(sp, "Admin")
    assert sp.may_write(users["admin"]) and not sp.may_write(operator)
    assert sp.write_access_text() == "Admin"
    limit_to(sp, "Operator")
    assert sp.may_write(operator) and not sp.may_write(users["admin"])  # applies to Admin too
    anon = AnonymousUser()
    assert not sp.may_write(anon)


def test_api_write_checks_roles(client, dashboards):
    """The write API refuses users whose role is not selected (403)."""
    make_user("op", "Operator")
    sp = dashboards["points"]["setpoint"]
    limit_to(sp, "Admin")
    login(client, "op")
    resp = client.post(f"/api/datapoints/{sp.id}/write", json={"value": 5})
    assert resp.status_code == 403 and "role" in resp.get_json()["error"]
    client.post("/auth/logout")
    login(client, "admin")
    assert client.post(f"/api/datapoints/{sp.id}/write", json={"value": 5}).status_code == 200


def test_grid_and_graphic_screens_show_controls_per_role(client, dashboards):
    """Write controls appear only for tags the user may write."""
    make_user("op", "Operator")
    sp, speed = dashboards["points"]["setpoint"], dashboards["points"]["speed"]
    limit_to(sp, "Admin")
    login(client, "op")
    html = client.get("/screens/private").get_data(as_text=True)
    assert f'data-id="{sp.id}"' in html
    tile = html.split(f'data-id="{sp.id}"')[1].split('data-id="')[0]
    assert 'data-writable="0"' in tile and "hmi-write" not in tile
    tile = html.split(f'data-id="{speed.id}"')[1].split('data-id="')[0]
    assert 'data-writable="1"' in tile and "hmi-write" in tile
    with client.application.test_request_context("/"):
        from flask_login import login_user  # noqa: PLC0415
        from app.dashboards.routes import _runtime_data  # noqa: PLC0415
        from app.models import User  # noqa: PLC0415

        login_user(User.query.filter_by(username="op").one())
        points = _runtime_data(dashboards["private"])["points"]
    assert points[str(sp.id)]["may_write"] is False and points[str(speed.id)]["may_write"] is True
    assert points[str(sp.id)]["write_roles"] == ["Admin"]


def test_tag_form_saves_roles(client, login_as, plc, db):
    """The tag form stores the access mode and the roles; roles are required for 'Based on user roles'."""
    login_as("admin")
    base = {**BASE, "plc_id": plc.id, "label": "Valve", "address": "DB3.DBD0", "writable": "y"}
    resp = client.post("/tags/new", data={**base, "write_access": "roles"})
    assert resp.status_code == 200 and "Select at least one role." in resp.get_data(as_text=True)
    op = role("Operator")
    resp = client.post("/tags/new", data={**base, "write_access": "roles", "write_roles": [op.id]})
    assert resp.status_code == 302, resp.get_data(as_text=True)
    tag = DataPoint.query.filter_by(label="Valve").one()
    assert tag.write_access == "roles" and [r.name for r in tag.write_roles] == ["Operator"]
    html = client.get(f"/tags/{tag.id}/edit").get_data(as_text=True)
    assert 'id="write-access-box"' in html and "write_access.js" in html
    box = html.split('id="write-roles-box"')[1].split("</div>\n</div>")[0]
    checked = re.findall(r'<input[^>]*name="write_roles"[^>]*>', box)
    assert any(f'value="{op.id}"' in i and "checked" in i for i in checked)
    assert not any(f'value="{op.id}"' not in i and "checked" in i for i in checked)
    assert "tag-write-roles" in client.get("/tags/").get_data(as_text=True)
    client.post(f"/tags/{tag.id}/edit", data={**base, "write_access": "all", "write_roles": [op.id]})
    tag = db.session.get(DataPoint, tag.id)
    assert tag.write_access == "all" and tag.write_roles == []


def test_screen_tag_form_has_write_access(client, login_as, dashboards):
    """The screen editor's tag form offers the same setting."""
    login_as("admin")
    sp = dashboards["points"]["setpoint"]
    html = client.get(f"/screens/{dashboards['private'].id}/points/{sp.id}/edit").get_data(as_text=True)
    assert 'name="write_access"' in html and 'name="write_roles"' in html


def test_excel_export_and_import_write_roles(client, login_as, dashboards, db):
    """'Write roles' is exported and imported (empty = everyone, unknown role = error)."""
    login_as("admin")
    sp = dashboards["points"]["setpoint"]
    limit_to(sp, "Operator", "Admin")
    wb = load_workbook(io.BytesIO(client.get("/tags/export").data))
    rows = list(wb["Tags"].iter_rows(values_only=True))
    col = rows[0].index("Write roles")
    row = next(r for r in rows[1:] if r[0] == sp.id)
    assert row[col] == "Admin, Operator"

    def upload(value):
        """Import one row changing the write roles of the set-point tag."""
        book = Workbook()
        sheet = book.active
        sheet.title = "Tags"
        sheet.append(["ID", "Label", "PLC", "Write roles"])
        sheet.append([sp.id, "Setpoint", "PLC1", value])
        buf = io.BytesIO()
        book.save(buf)
        buf.seek(0)
        return client.post("/tags/import", data={"file": (buf, "t.xlsx")}, content_type="multipart/form-data")

    assert "Unknown role(s): Boss" in upload("Boss").get_data(as_text=True)
    upload("Operator")
    tag = db.session.get(DataPoint, sp.id)
    assert tag.write_access == "roles" and [r.name for r in tag.write_roles] == ["Operator"]
    upload(None)
    tag = db.session.get(DataPoint, sp.id)
    assert tag.write_access == "all" and tag.write_roles == []


def test_deleting_a_role_removes_it_from_tags(client, login_as, dashboards, db):
    """A deleted custom role disappears from the tags it was allowed to write."""
    custom = Role(name="Shift lead", description="")
    db.session.add(custom)
    db.session.commit()
    sp = dashboards["points"]["setpoint"]
    limit_to(sp, "Shift lead", "Admin")
    login_as("admin")
    client.post(f"/admin/roles/{custom.id}/delete")
    assert Role.query.filter_by(name="Shift lead").count() == 0
    assert [r.name for r in db.session.get(DataPoint, sp.id).write_roles] == ["Admin"]


def test_details_show_write_access(client, login_as, dashboards):
    """The Details window lists who may write and whether the user may."""
    sp = dashboards["points"]["setpoint"]
    limit_to(sp, "Operator")
    login_as("admin")
    d = client.get(f"/api/datapoints/{sp.id}/details").get_json()
    assert d["write_roles"] == ["Operator"] and d["may_write"] is False


def test_local_dummy_tag_respects_roles(app, db, users):
    """Local dummy tags need no write permission, but a role limit still applies."""
    from app.models import PLC  # noqa: PLC0415
    from app.models.user import AnonymousUser  # noqa: PLC0415

    dummy = PLC(name="Internal", protocol="dummy", ip_address="", rack=0, slot=0)
    db.session.add(dummy)
    db.session.commit()
    from app.plc.dummy import DummyAddress  # noqa: PLC0415

    local = DataPoint(plc=dummy, label="Local", writable=True, dummy_scope="local")
    local.set_address(DummyAddress(start=1, data_type="INT"))
    db.session.add(local)
    db.session.commit()
    assert local.is_local and local.may_write(AnonymousUser()) and local.may_write(users["viewer"])
    limit_to(local, "Admin")
    assert local.may_write(users["admin"]) and not local.may_write(users["viewer"])


def test_write_access_changes_are_audited(client, login_as, plc):
    """Changing who may write a tag is recorded in the audit log with the role names."""
    from app.models import AuditLog  # noqa: PLC0415

    login_as("admin")
    base = {**BASE, "plc_id": plc.id, "label": "Pump", "address": "DB3.DBD0", "writable": "y"}
    client.post("/tags/new", data={**base, "write_access": "all"})
    tag = DataPoint.query.filter_by(label="Pump").one()
    client.post(f"/tags/{tag.id}/edit", data={**base, "write_access": "roles", "write_roles": [role("Operator").id]})
    entry = AuditLog.query.filter_by(action="update", object_type="Tag", object_id=tag.id).order_by(AuditLog.id.desc()).first()
    changes = {c["field"]: c for c in entry.changes}
    assert changes["write_access"]["after"] == "roles"
    assert changes["write_roles"]["after"] == "Operator"
